"""Read the supplied plan/KRM without editing the workbooks or student data."""
import argparse,hashlib,json,re
from pathlib import Path
import openpyxl

def extract(plan_path,krm_path,registry_path):
    registry=json.loads(Path(registry_path).read_text(encoding='utf-8-sig'))
    target={x['code'] for x in registry['competencies']}
    krm=openpyxl.load_workbook(krm_path,read_only=True,data_only=True)
    cells={code:[] for code in target}
    for sheet in krm:
        for row in sheet:
            for cell in row:
                value=str(cell.value or '').strip()
                if value in cells and len(cells[value])<3:cells[value].append({'sheet':sheet.title,'cell':cell.coordinate})
    krm.close()
    plan=openpyxl.load_workbook(plan_path,read_only=False,data_only=True)
    sheet=plan['Дисциплины'];courses=[];modules=[]
    columns=[('X','Y','Z'),('AB','AC','AD'),('AG','AH','AI'),('AK','AL','AM'),('AP','AQ','AR'),('AT','AU','AV'),('AY','AZ','BA'),('BC','BD','BE')]
    for row in range(20,sheet.max_row+1):
        code=sheet.cell(row,1).value;title=sheet.cell(row,2).value;raw=sheet.cell(row,4).value
        if code and title and re.match(r'^Б\.1\.2\.2\.2\.\d+$',str(code)):
            modules.append({'id':str(code).split('.')[-1],'title':str(title).strip(),'sourceCell':f'B{row}'})
        if not code or not title or not raw:continue
        codes=re.findall(r'[A-ZА-Я]+(?:-[A-ZА-Я]+)*-\d+',str(raw))
        known=[c for c in codes if c in target and cells.get(c)]
        semesters=[]
        for semester,cols in enumerate(columns,1):
            if any(isinstance(sheet[f'{col}{row}'].value,(int,float)) and sheet[f'{col}{row}'].value>0 for col in cols):semesters.append(semester)
        courses.append({'planCode':str(code),'title':str(title).strip(),'row':row,'competencyCell':f'D{row}','planCompetencies':codes,'krmCompetencies':known,'unresolvedPlanCodes':[c for c in codes if c not in target],'semesters':semesters,'credits':sheet.cell(row,16).value})
    plan.close()
    def source(path):
        p=Path(path);return {'filename':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
    return {'version':'pmo-curriculum-2026-10-29','sources':{'plan':source(plan_path),'krm':source(krm_path),'targetModelRegistry':source(registry_path)},'krmCompetencySources':cells,'electiveModules':modules,'electiveRequired':2,'courses':courses,'note':'КРМ подтверждает наличие компетенций, а учебный план – их связь с дисциплиной. Связь конкретного задания с индикатором является отдельным методическим предложением.'}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--plan',required=True);p.add_argument('--krm',required=True);p.add_argument('--registry',required=True);p.add_argument('--output',required=True);args=p.parse_args()
    data=extract(args.plan,args.krm,args.registry);Path(args.output).write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    for c in data['courses']:print(c['row'],c['title'],','.join(c['krmCompetencies']) or '[нет кода целевой модели]',c['semesters'])
