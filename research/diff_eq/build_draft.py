"""Build editable research draft and self-contained LaTeX from shared content."""
from pathlib import Path
import sys
BASE=Path(__file__).resolve().parent
if (BASE/'deps').exists():sys.path.insert(0,str(BASE/'deps'))
from manuscript import TITLE,ABSTRACT,BLOCKS,REFERENCES,R
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from latex2mathml.converter import convert
from zipfile import ZipFile
import json,re

import os
OUT=Path(os.environ.get('DOCTOR_PIT_OUTPUT_DIR',str(BASE.parent.parent/'outputs')))
OUT.mkdir(parents=True,exist_ok=True)
STEM='Модель_ОДУ_Идентифицируемость_рабочая_статья'
OUT.mkdir(exist_ok=True)
XSL=etree.XSLT(etree.parse('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
doc=Document()
sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(2);sec.bottom_margin=Cm(2)
sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)
sec.header_distance=Cm(.8);sec.footer_distance=Cm(.8)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11.5)
normal.font.color.rgb=RGBColor(0,0,0)
normal.paragraph_format.line_spacing=1.05
normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.first_line_indent=Cm(.65)
normal.paragraph_format.widow_control=True
for name in ['Title','Subtitle','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.first_line_indent=Cm(0)
    for tag in ['w:pBdr','w:shd']:
        for el in list(st.element.iter(qn(tag))):el.getparent().remove(el)
doc.styles['Title'].font.size=Pt(16)
doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(9)
doc.styles['Heading 1'].font.size=Pt(12)
doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(5)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10)
doc.styles['Caption'].paragraph_format.space_after=Pt(5)
doc.styles['Caption'].paragraph_format.keep_with_next=True

def text(txt,style=None,indent=True,align=None):
    p=doc.add_paragraph(txt,style)
    if not indent:p.paragraph_format.first_line_indent=Cm(0)
    if align is not None:p.alignment=align
    elif style not in ['Title','Heading 1','Caption']:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

text(TITLE,'Title',False,WD_ALIGN_PARAGRAPH.CENTER)
p=text('П. В. Никитин',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER);p.runs[0].bold=True
p=text('Финансовый университет при Правительстве Российской Федерации\nМосква Россия\npvnikitin@fa.ru   ORCID 0000-0001-8866-5610',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs:r.font.size=Pt(10)
p=text('Рабочая математическая редакция от 22 октября 2026 года',indent=False,align=WD_ALIGN_PARAGRAPH.CENTER)
for r in p.runs:r.italic=True;r.font.size=Pt(10)
p=text('Аннотация. '+ABSTRACT,indent=False)
for r in p.runs:r.font.size=Pt(10.5)
p=text('Ключевые слова: нелинейные ОДУ, структурная идентифицируемость, конечные наблюдения, междисциплинарный перенос, внешняя поддержка, обратная задача.',indent=False)
for r in p.runs:r.font.size=Pt(10)

eq_count=0
def equation(tex):
    global eq_count
    eq_count+=1
    math=etree.fromstring(convert(tex).encode())
    omml=XSL(math).getroot()
    if omml.tag==qn('m:oMathPara'):
        omml=omml.find(qn('m:oMath'))
    for t in omml.findall('.//'+qn('m:r')):
        rp=t.find(qn('w:rPr'))
        if rp is None:rp=OxmlElement('w:rPr');t.insert(0,rp)
        fonts=OxmlElement('w:rFonts');fonts.set(qn('w:ascii'),'Cambria Math');fonts.set(qn('w:hAnsi'),'Cambria Math');rp.append(fonts)
        size=OxmlElement('w:sz');size.set(qn('w:val'),'21');rp.append(size)
    p=doc.add_paragraph()
    p.paragraph_format.first_line_indent=Cm(0)
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.keep_together=True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(8.0),WD_ALIGN_PARAGRAPH.CENTER)
    # Explicit Word tab alignment values are used below, independent of paragraph enums.
    tabs=p._p.get_or_add_pPr().find(qn('w:tabs'))
    for child in list(tabs):tabs.remove(child)
    for val,pos in [('center','4450'),('right','9400')]:
        tab=OxmlElement('w:tab');tab.set(qn('w:val'),val);tab.set(qn('w:pos'),pos);tabs.append(tab)
    p.add_run('\t');p._p.append(omml);p.add_run('\t('+str(eq_count)+')')

def make_table(info):
    text(info['caption'],'Caption',False)
    table=doc.add_table(rows=1,cols=len(info['heads']))
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    widths=([2.7,3.6,10.3] if len(info['heads'])==3 else [1.2,1.7,5.4,8.3])
    if len(info['heads'])==4 and info['heads'][2].startswith('Калибровка'):widths=[1.6,2.1,6.2,6.7]
    for i,v in enumerate(info['heads']):table.rows[0].cells[i].text=v
    for row in info['rows']:
        for cell,v in zip(table.add_row().cells,row):cell.text=v
    for ri,row in enumerate(table.rows):
        trpr=row._tr.get_or_add_trPr()
        if ri==0:
            repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');trpr.append(repeat)
        no_split=OxmlElement('w:cantSplit');trpr.append(no_split)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            cp.append(borders)
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'70');el.set(qn('w:type'),'dxa');margins.append(el)
            cp.append(margins)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6E6E6');cp.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(widths)==3 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:r.font.size=Pt(10);r.bold=ri==0
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

for kind,value in BLOCKS:
    if kind=='h':text(value,'Heading 1',False)
    elif kind=='p':text(value)
    elif kind=='s':
        p=text(value)
        for r in p.runs:r.italic=True
    elif kind=='e':equation(value)
    elif kind=='t':make_table(value)

text('Литература','Heading 1',False)
for i,ref in enumerate(REFERENCES,1):
    p=text(str(i)+'. '+ref,indent=False)
    p.paragraph_format.left_indent=Cm(.5);p.paragraph_format.first_line_indent=Cm(-.5)
    for r in p.runs:r.font.size=Pt(10)
text('Identifiability of a Two Component Model of Interdisciplinary Learning under Switching Assessment Modes','Heading 1',False)
p=text('P. V. Nikitin\nFinancial University under the Government of the Russian Federation Moscow Russia\npvnikitin@fa.ru   ORCID 0000-0001-8866-5610',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT)
for r in p.runs:r.font.size=Pt(10)
p=text('The draft studies a triangular nonlinear ODE with learning, forgetting and state dependent transfer. '
       'An observation mechanism separates independent performance from externally supported performance. '
       'A finite protocol identifies nine unknown parameters from nine exact success probabilities. '
       'Global injectivity and local inverse stability are proved under explicit assumptions. '
       'All computational results are synthetic. The validity of the educational interpretation remains to be tested.',indent=False)
for r in p.runs:r.font.size=Pt(10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=footer.add_run('Рабочая редакция   ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title=TITLE;doc.core_properties.author='Пётр Владимирович Никитин'
doc.core_properties.subject='Рабочая математическая статья с доказательствами и синтетической проверкой'
doc.save(OUT/(STEM+'.docx'))

def esc(txt):
    # Body math notation is rendered using Unicode; use LuaLaTeX for this standalone source.
    return txt.replace('\\',r'\textbackslash{}').replace('&',r'\&').replace('%',r'\%').replace('#',r'\#').replace('_',r'\_')

tex=[r'\documentclass[11pt,a4paper]{article}',
     r'\usepackage{fontspec}',r'\setmainfont{Times New Roman}',
     r'\usepackage{unicode-math}',r'\setmathfont{Cambria Math}',
     r'\usepackage[russian,english]{babel}',r'\usepackage{amsmath}',
     r'\usepackage[margin=22mm]{geometry}',r'\usepackage{longtable,array}',
     r'\setlength{\parindent}{1.2em}',r'\setlength{\parskip}{3pt}',
     r'\begin{document}',r'\selectlanguage{russian}',
     r'\begin{center}{\Large\bfseries '+esc(TITLE)+r'}\par\medskip',
     r'П. В. Никитин\par Финансовый университет при Правительстве Российской Федерации\par Москва Россия\par',
     r'\texttt{pvnikitin@fa.ru}\quad ORCID 0000-0001-8866-5610\par',
     r'\small Рабочая математическая редакция от 22 октября 2026 года\end{center}',
     r'\noindent\textbf{Аннотация.} '+esc(ABSTRACT)+r'\par',
     r'\noindent\textbf{Ключевые слова:} нелинейные ОДУ, идентифицируемость, конечные наблюдения, перенос, поддержка.']
for kind,value in BLOCKS:
    if kind=='h':tex.append(r'\section*{'+esc(value)+'}')
    elif kind=='e':tex.append(r'\begin{equation}'+value+r'\end{equation}')
    elif kind=='s':tex.append(r'\textit{'+esc(value)+r'}\par')
    elif kind=='p':tex.append(esc(value)+r'\par')
    elif kind=='t':
        cols=len(value['heads'])
        fmt='|'+'|'.join(['p{'+('%.3f'%(0.90/cols))+r'\linewidth}' for _ in range(cols)])+'|'
        tex.append(r'\begin{center}\small '+esc(value['caption'])+r'\end{center}')
        tex.append(r'\begin{longtable}{'+fmt+r'}\hline')
        tex.append(' & '.join(esc(v) for v in value['heads'])+r'\\\hline\endhead')
        for row in value['rows']:tex.append(' & '.join(esc(v) for v in row)+r'\\\hline')
        tex.append(r'\end{longtable}')
tex.append(r'\section*{Литература}\begin{enumerate}')
for ref in REFERENCES:tex.append(r'\item '+esc(ref))
tex.extend([r'\end{enumerate}',r'\selectlanguage{english}',
    r'\section*{Identifiability of a Two Component Model of Interdisciplinary Learning under Switching Assessment Modes}',
    r'P. V. Nikitin\par Financial University under the Government of the Russian Federation Moscow Russia\par',
    r'\texttt{pvnikitin@fa.ru}\quad ORCID 0000-0001-8866-5610',r'\end{document}'])
(OUT/(STEM+'.tex')).write_text('\n\n'.join(tex),encoding='utf-8')
(BASE/'manuscript.json').write_text(json.dumps(dict(title=TITLE,abstract=ABSTRACT,blocks=BLOCKS,references=REFERENCES),ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(docx=str(OUT/(STEM+'.docx')),tex=str(OUT/(STEM+'.tex')),equations=eq_count,
                     tables=3,words=sum(len(str(v).split()) for k,v in BLOCKS)),ensure_ascii=False,indent=2))
