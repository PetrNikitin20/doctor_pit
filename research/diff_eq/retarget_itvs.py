"""Adapt the mathematical manuscript to ITVS author instructions.

The model and computed results are preserved. Submission is not performed.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree as E
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json,re
from manuscript import BLOCKS,R,REFERENCES,_shared_notation

B=Path(__file__).resolve().parent
import os
OUT=Path(os.environ.get('DOCTOR_PIT_OUTPUT_DIR',str(B.parent.parent/'outputs')))
OUT.mkdir(parents=True,exist_ok=True)
SRC=OUT/'Модель_ОДУ_Идентифицируемость_рабочая_статья.docx'
FINAL=OUT/'Статья_ИТВС_Идентификация_динамической_модели_подготовки.docx'
TITLE='Идентификация динамической модели междисциплинарной подготовки при внешней поддержке'
TITLE_EN='Identifying a Dynamic Interdisciplinary Learning Model with External Assistance'
ABSTRACT=('Рассматривается задача восстановления самостоятельной подготовки по результатам заданий, '
'выполняемых самостоятельно и с внешней поддержкой. Предложена двухкомпонентная динамическая '
'модель освоения, забывания и направленного междисциплинарного переноса. Неизвестное влияние '
'поддержки учитывается в механизме наблюдения. Построен конечный протокол идентификации '
'девяти параметров по девяти точным вероятностям успеха. Доказаны однозначность восстановления '
'и локальная устойчивость обратного отображения. Сопоставлены смешанный и полностью '
'поддерживаемый режимы диагностики. Выполнены восемьдесят восстановлений на синтетических '
'данных и статистическое моделирование ошибок калибровки. Показано ухудшение обусловленности '
'при коротком интервале наблюдения и почти постоянной предпосылке. Результат позволяет '
'выявлять условия информативности диагностики в интеллектуальной обучающей системе. '
'Применимость модели к реальным обучающимся требует отдельной проверки.')
ABSTRACT_EN=('We study the recovery of independent preparation from tasks completed independently '
'or with external assistance. A two component dynamical model describes learning, forgetting '
'and directed interdisciplinary transfer. An unknown assistance effect enters the observation '
'mechanism. A finite protocol identifies nine parameters from nine exact success probabilities. '
'Uniqueness and local inverse stability are proved. Mixed and fully assisted assessment regimes '
'are compared through eighty synthetic recovery cases and statistical calibration experiments. '
'Short observation intervals and nearly constant prerequisite states impair conditioning. '
'The results characterize informative assessment conditions for an intelligent learning system. '
'The educational interpretation and applicability to real learners require separate validation.')
KEYS='идентифицируемость, динамическая модель, междисциплинарное обучение, внешняя поддержка, диагностика, численный эксперимент, интеллектуальная обучающая система'
KEYS_EN='identifiability, dynamical model, interdisciplinary learning, external assistance, assessment, numerical experiment, intelligent learning system'

sections={}
current=None
for kind,value in BLOCKS:
    if kind=='h':current=int(value.split()[0]);sections[current]=[]
    else:sections[current].append((kind,value))

new=[]
def p(t):new.append(('p',t))
def h(t):new.append(('h',t))
def add(old):new.extend(deepcopy(sections[old]))

h('Введение')
add(1)
for i,(kind,value) in enumerate(new):
    if kind=='p' and 'Приоритет полученных' in value:
        new[i]=(kind,value[:value.index(' Приоритет полученных')]+ ' Вклад работы состоит в конструктивном протоколе '
                'и доказательстве однозначности восстановления для заданной двухкомпонентной системы с неизвестной поддержкой.')
p('Цель работы – установить условия, при которых конечные наблюдения разделяют изменение самостоятельной '
  'подготовки, перенос между компонентами и влияние внешней помощи. Сначала задаются модель и протокол, '
  'затем обосновываются восстановление и его устойчивость. После описания вычислительного алгоритма '
  'приводятся синтетическая проверка и границы применения результата.')

h('1 Модель подготовки и механизм наблюдений')
add(2)
h('2 Протокол и конструктивное восстановление параметров')
add(3)
add(4)
p('В полностью поддерживаемом режиме восстановление постоянного слагаемого и одной экспоненты '
  'является частным случаем классической задачи оценивания экспоненциальных сумм [8]. '
  'Эти алгебраические формулы не предлагаются как новый вариант метода Прони. Их роль здесь – '
  'калибровка неизвестной поддержки перед идентификацией нелинейного междисциплинарного взаимодействия.')
add(5)
h('3 Идентифицируемость и устойчивость восстановления')
add(6)
add(7)
add(8)

h('4 Вычислительный алгоритм и исследовательская реализация')
p('Исследовательская программа реализует последовательную идентификацию, следующую из теоремы 1. '
  'На вход подаются девять вероятностей и известное расписание учебных воздействий. Для данных '
  'в допустимом образе сначала вычисляются параметры калибровки двух компонентов. Затем из '
  'седьмого наблюдения определяется состояние предпосылки и решается скалярное уравнение для α₁. '
  'На последнем шаге два оставшихся параметра определяются по двум наблюдениям второго компонента.')
p('Скалярное уравнение решается методом Брента после проверки знаков на концах интервала. '
  'Для пары (α₂,β) применяется ограниченный метод наименьших квадратов с аналитическим '
  'якобианом чувствительностей. Используются стандартные средства SciPy [9]. Если прямые '
  'формулы калибровки выходят за область параметров, программа фиксирует отказ такого обращения '
  'и сохраняет исходные частоты. Автоматическое обрезание параметров до допустимых границ '
  'изменило бы обратную задачу и не используется в данном эксперименте.')
p('Прямая задача на последнем участке решается по формуле вариации постоянных. В обозначениях '
  'леммы 2 вводятся K(t) = ∫₀ᵗ[c₂(α₂ + βw(s)) + λ₂]ds и z(0) = B. Тогда '
  '1 − z(t) = (1 − B)exp(−K(t)) + λ₂∫₀ᵗexp(K(s) − K(t))ds. '
  'Запись через разность K(s) − K(t) позволяет избежать положительных больших экспонент. '
  'Интеграл вычисляется квадратурой с 64 узлами. Независимый контроль использует интегрирование '
  'исходной системы, а не ту же интегральную формулу.')
p('При фиксированном числе узлов стоимость прямого расчёта и аналитических чувствительностей '
  'пропорциональна числу запрошенных моментов. Стоимость обратного расчёта дополнительно зависит '
  'от числа итераций методов поиска корня и наименьших квадратов. Глобальная сходимость '
  'оптимизатора из произвольной начальной точки из теоремы об идентифицируемости не следует. '
  'Поэтому сохраняются начальная точка, ограничения, невязка и признак завершения расчёта.')
p('Реализован вычислительный исследовательский модуль. Его связь с учебным журналом задаётся '
  'через записи о компоненте, классе задания, времени, учебном воздействии и режиме поддержки. '
  'Подключение модуля к рабочему образовательному комплексу и проверка на реальных событиях '
  'рассматриваются как следующий этап. Входные вероятности нельзя непосредственно заменить '
  'оценкой за одно задание или усреднить между разными студентами без модели неоднородности.')

h('5 Вычислительный эксперимент')
add(9)
p('Таблица 2 сопоставляет обусловленность калибровки и полного отображения наблюдений; '
  'таблица 3 – чувствительность прямых формул к биномиальному шуму. При h = 0,1 и n = 10000 '
  'доля троек вне допустимого образа составила 1,88 % для M и 57,98 % для A. '
  'При h = 3 в этой конфигурации обе доли равны нулю. Эти результаты не устанавливают '
  'оптимальное расписание для всех параметров или моделей ошибок.')
p('На рисунке 1 приведены две зависимости: уменьшение информативности короткой калибровки '
  'и вырождение разделения освоения и переноса при почти постоянной предпосылке. '
  'Нормировка параметров фиксирована и совпадает с нормировкой таблицы 2.')
new.append(('f',dict(file='conditioning.png',caption='Рис. 1. Обусловленность калибровки и идентификации переноса')))

h('6 Применение и ограничения')
new.extend(deepcopy(sections[10][:-1]))
h('Заключение')
p('Для двухкомпонентной модели междисциплинарной подготовки построен конечный протокол '
  'восстановления девяти неизвестных параметров, включая характеристики внешней поддержки. '
  'Доказаны глобальная однозначность восстановления на образе модели, невырожденность якобиана '
  'и локальная устойчивость обратного отображения. Контрпример постоянной предпосылки '
  'указывает на необходимость динамически информативного расписания.')
p('Исследовательская реализация проверена на 80 синтетических случаях восстановления '
  'и 180000 статистических повторениях калибровки. Сравнение режимов M и A показывает '
  'роль независимых диагностических наблюдений при коротком интервале. Следующие задачи – '
  'проверка наблюдательного механизма по индивидуальным учебным событиям и расширение '
  'идентификационного протокола на несколько взаимосвязанных компонентов.')

# Normalize notation in the application paragraphs added for this journal too.
new=[(kind,_shared_notation(value) if kind in ['p','s','e'] and not value.startswith('Вектор неизвестных параметров обозначен') else value) for kind,value in new]

# Renumber references by first appearance after inserting the classical calibration source.
first=[]
for kind,value in new:
    if kind in ['p','s']:
        for v in re.findall(r'\[(\d+)\]',value):
            n=int(v)
            if n not in first:first.append(n)
mapping={old:i+1 for i,old in enumerate(first)}
rawrefs={i+1:t for i,t in enumerate(REFERENCES)}
rawrefs[8]='Potts D., Tasche M. Parameter estimation for exponential sums by approximate Prony method // Signal Processing. 2010. V. 90. № 5. P. 1631–1642. DOI: 10.1016/j.sigpro.2009.11.012.'
rawrefs[9]='Virtanen P., Gommers R., Oliphant T. E. et al. SciPy 1.0: fundamental algorithms for scientific computing in Python // Nature Methods. 2020. V. 17. № 3. P. 261–272. DOI: 10.1038/s41592-019-0686-2.'
rawrefs[1]=rawrefs[1].replace(' Год приведён по странице издателя Springer.','')
rawrefs[6]=rawrefs[6][:rawrefs[6].index(' Рабочая ссылка')]
refs=[rawrefs[i] for i in first]
for i,(kind,value) in enumerate(new):
    if kind in ['p','s']:
        value=re.sub(r'\[(\d+)\]',lambda m:'['+str(mapping[int(m[1])])+']',value)
        value=value.replace('условия протокола раздела 3','условия протокола раздела 2')
        value=value.replace('указанного в разделе 2','указанного в разделе 1')
        if 'Обозначим T₀' in value:value=value.replace('В таблице приведён','В таблице 1 приведён')
        if 'Рассматривается система' in value:value=value.replace('Рассматривается система','Рассматривается система (1), (2)')
        if 'Перенос действует только' in value:value+=' Вероятность наблюдаемого успеха задаётся формулой (3).'
        if 'При фиксированных внутренних значениях' in value:value='Определители (4) дают следующее сравнение. '+value
        if value.endswith('Поэтому'):value+=' определитель полного отображения имеет вид (5).'
        if 'Если оценка ψ̂' in value:value='Из оценки (6) следует, что '+value[0].lower()+value[1:]
        new[i]=(kind,value)

# Native editable formulas from the verified source draft.
with ZipFile(SRC) as z:root=E.fromstring(z.read('word/document.xml'))
maths=root.findall('.//'+qn('m:oMath'))
assert len(maths)==22
numbered={1:1,2:2,3:3,10:4,19:5,20:6}
doc=Document();sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(3.6);sec.bottom_margin=Cm(3.6)
sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
sec.header_distance=Cm(1.5);sec.footer_distance=Cm(1.5)
for st in doc.styles:
    if st.type==1:
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        pr=st.element.get_or_add_rPr()
        ff=pr.find(qn('w:rFonts'))
        if ff is None:ff=OxmlElement('w:rFonts');pr.append(ff)
        for key in list(ff.attrib):
            if 'Theme' in key:del ff.attrib[key]
        ff.set(qn('w:ascii'),'Times New Roman');ff.set(qn('w:hAnsi'),'Times New Roman')
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.5
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.first_line_indent=Cm(.5)
normal.paragraph_format.widow_control=True
for name in ['Title','Heading 1','Caption']:
    st=doc.styles[name];st.font.color.rgb=RGBColor(0,0,0)
    st.paragraph_format.first_line_indent=Cm(0)
    for tag in ['w:pBdr','w:shd']:
        for el in list(st.element.iter(qn(tag))):el.getparent().remove(el)
doc.styles['Title'].font.size=Pt(12);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 1'].font.size=Pt(11);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(9)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(4)
doc.styles['Heading 1'].paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(10)

def para(t,style=None,indent=True,align=None,single=False,size=None):
    p=doc.add_paragraph(t,style)
    p.alignment=align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    if not indent:p.paragraph_format.first_line_indent=Cm(0)
    if single:p.paragraph_format.line_spacing=1
    if size:
        for r in p.runs:r.font.size=Pt(size)
    return p

para(TITLE,'Title',False,WD_ALIGN_PARAGRAPH.CENTER)
para('П. В. Никитин',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT)
para('Финансовый университет при Правительстве Российской Федерации, Москва, Россия',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT)
para('Аннотация. '+ABSTRACT,indent=False,single=True)
para('Ключевые слова: '+KEYS+'.',indent=False,single=True)

en_refs={
1:'Corbett AT, Anderson JR. Knowledge tracing: Modeling the acquisition of procedural knowledge. User Modeling and User-Adapted Interaction. 1994;4:253–278. doi:10.1007/BF01099821.',
2:'Bellman R, Åström KJ. On structural identifiability. Mathematical Biosciences. 1970;7(3–4):329–339. doi:10.1016/0025-5564(70)90132-X.',
3:'Miao H, Xia X, Perelson AS, Wu H. On Identifiability of Nonlinear ODE Models and Applications in Viral Dynamics. SIAM Review. 2011;53(1):3–39. doi:10.1137/090757009.',
4:'Simpson MJ, Baker RE. Parameter Identifiability, Parameter Estimation, and Model Prediction for Differential Equation Models. SIAM Review. 2026;68(1):153–171. doi:10.1137/24M1667968.',
5:'Wang Y, Huang W, Gong M, Geng X, Liu T, Zhang K, Tao D. Identifiability and Asymptotics in Learning Homogeneous Linear ODE Systems from Discrete Observations. Journal of Machine Learning Research. 2024;25(154):1–50. Available from: https://jmlr.org/papers/v25/22-1159.html.',
6:'Kostanian AA, Beklaryan AL. Identifiability of a dissipative knowledge-dynamics model: exact recovery under designed excitation, degeneration on observational data. Preprint. arXiv:2610.09889v3. 2026 Oct 15. doi:10.48550/arXiv.2610.09889.',
7:'Hoeffding W. Probability Inequalities for Sums of Bounded Random Variables. Journal of the American Statistical Association. 1963;58(301):13–30. doi:10.1080/01621459.1963.10500830.',
8:'Potts D, Tasche M. Parameter estimation for exponential sums by approximate Prony method. Signal Processing. 2010;90(5):1631–1642. doi:10.1016/j.sigpro.2009.11.012.',
9:'Virtanen P, Gommers R, Oliphant TE, et al. SciPy 1.0: fundamental algorithms for scientific computing in Python. Nature Methods. 2020;17(3):261–272. doi:10.1038/s41592-019-0686-2.'}

def tbl(info):
    p=para(info['caption'],'Caption',False,single=True,size=10)
    p.paragraph_format.keep_with_next=True
    tb=doc.add_table(rows=1,cols=len(info['heads']));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False
    widths=[2.5,3.1,10.4] if len(info['heads'])==3 else [1.2,1.8,5.4,7.6]
    for c,w in zip(tb.columns,widths):c.width=Cm(w)
    for cell,v in zip(tb.rows[0].cells,info['heads']):cell.text=v
    for row in info['rows']:
        for cell,v in zip(tb.add_row().cells,row):cell.text=v.replace('e-','e−')
    for ri,row in enumerate(tb.rows):
        trpr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');trpr.append(cs)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
            pr.append(bd)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E6E6E6');pr.append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1
                p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
                if ri==0:p.paragraph_format.keep_with_next=True
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if len(widths)==3 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:run.font.size=Pt(10.5);run.bold=ri==0

eqindex=0
previous_kind=None
for block_index,(kind,value) in enumerate(new):
    if kind=='h':para(value,'Heading 1',False,WD_ALIGN_PARAGRAPH.LEFT)
    elif kind in ['p','s']:
        p=para(value)
        if previous_kind=='t':p.paragraph_format.space_before=Pt(6)
        if kind=='s':
            for run in p.runs:run.italic=True
    elif kind=='t':tbl(value)
    elif kind=='e':
        eqindex+=1;p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(5)
        p.paragraph_format.keep_together=True
        if block_index+1<len(new) and new[block_index+1][0]=='e':
            p.paragraph_format.keep_with_next=True
        tabs=OxmlElement('w:tabs')
        for val,pos in [('center','4200'),('right','9070')]:
            tab=OxmlElement('w:tab');tab.set(qn('w:val'),val);tab.set(qn('w:pos'),pos);tabs.append(tab)
        p._p.get_or_add_pPr().append(tabs);p.add_run('\t');p._p.append(deepcopy(maths[eqindex-1]))
        if eqindex in numbered:p.add_run('\t('+str(numbered[eqindex])+')')
    elif kind=='f':
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_with_next=True
        p.add_run().add_picture(str(B/value['file']),width=Cm(16))
        p=para(value['caption'],'Caption',False,WD_ALIGN_PARAGRAPH.CENTER,single=True)
        p.paragraph_format.keep_with_next=False
    previous_kind=kind

para('Литература','Heading 1',False,WD_ALIGN_PARAGRAPH.LEFT)
for i,t in enumerate(refs,1):para(str(i)+'. '+t,indent=False,single=True,size=10)
para('Об авторе','Heading 1',False,WD_ALIGN_PARAGRAPH.LEFT)
para('Никитин Петр Владимирович. Кандидат педагогических наук, доцент. Доцент кафедры искусственного '
     'интеллекта Финансового университета при Правительстве Российской Федерации, Москва, Россия. '
     'Научные интересы: математическое моделирование образовательных процессов, искусственный интеллект '
     'и интеллектуальные обучающие системы. E-mail: pvnikitin@fa.ru (ответственный за переписку). '
     'ORCID: 0000-0001-8866-5610.',indent=False,single=True,size=10)
para(TITLE_EN,'Title',False,WD_ALIGN_PARAGRAPH.CENTER)
para('P. V. Nikitin',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT,single=True)
para('Financial University under the Government of the Russian Federation, Moscow, Russia',indent=False,align=WD_ALIGN_PARAGRAPH.LEFT,single=True)
para('Abstract. '+ABSTRACT_EN,indent=False,single=True)
para('Keywords: '+KEYS_EN+'.',indent=False,single=True)
para('References','Heading 1',False,WD_ALIGN_PARAGRAPH.LEFT)
for i,old in enumerate(first,1):para(str(i)+'. '+en_refs[old],indent=False,single=True,size=10)
para('About the author','Heading 1',False,WD_ALIGN_PARAGRAPH.LEFT)
para('Petr V. Nikitin. Candidate of Pedagogical Sciences, Associate Professor. Associate Professor, '
     'Department of Artificial Intelligence, Financial University under the Government of the Russian Federation, '
     'Moscow, Russia. Research interests: mathematical modeling of educational processes, artificial intelligence '
     'and intelligent learning systems. E-mail: pvnikitin@fa.ru (corresponding author). '
     'ORCID: 0000-0001-8866-5610.',indent=False,single=True,size=10)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
doc.core_properties.title=TITLE;doc.core_properties.author='Пётр Владимирович Никитин'
doc.core_properties.subject='Рукопись для журнала Информационные технологии и вычислительные системы'
doc.save(FINAL)
content=dict(title=TITLE,title_en=TITLE_EN,abstract=ABSTRACT,abstract_en=ABSTRACT_EN,keywords=KEYS,
             blocks=new,references=refs,references_en=[en_refs[i] for i in first],reference_mapping=mapping,
             numbered_equations=numbered,journal='Информационные технологии и вычислительные системы')
(B/'manuscript_itvs.json').write_text(json.dumps(content,ensure_ascii=False,indent=2),encoding='utf8')
summary=dict(file=str(FINAL),title_words_ru=len(TITLE.split()),title_words_en=len(TITLE_EN.split()),
             abstract_words_ru=len(ABSTRACT.split()),abstract_words_en=len(ABSTRACT_EN.split()),
             characters_with_spaces=sum(len(p.text) for p in doc.paragraphs)+sum(len(c.text) for t in doc.tables for row in t.rows for c in row.cells),
             equations=eqindex,numbered_equations=len(numbered),references=len(refs),figures=len(doc.inline_shapes))
(B/'retarget_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(summary,ensure_ascii=False,indent=2))
