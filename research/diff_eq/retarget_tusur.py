"""Build the TUSUR working manuscript from the unchanged ITVS model.

Requires an official publisher template, python-docx, latex2mathml, Office's
MML2OMML transform and the pinned scientific dependencies for the figure.
This is an author-review document, not a submission or MathType conversion.
"""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
import argparse, hashlib, json, re, tempfile
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from latex2mathml.converter import convert

B = Path(__file__).resolve().parent
TITLE = 'Идентификация динамической модели междисциплинарной подготовки при внешней поддержке'
TITLE_EN = 'Identifying a Dynamic Interdisciplinary Learning Model with External Assistance'
ABSTRACT = [
 ('Актуальность.', 'Результаты заданий с внешней помощью могут завышать оценку самостоятельной подготовки, поэтому требуется разделять динамику состояния и влияние поддержки.'),
 ('Цель исследования.', 'Определить условия однозначного восстановления параметров двухкомпонентной модели освоения, забывания и направленного междисциплинарного переноса.'),
 ('Методы.', 'Использованы анализ обыкновенных дифференциальных уравнений, конструктивное обращение конечного отображения наблюдений, уравнения чувствительности, анализ сингулярных чисел и биномиальное моделирование.'),
 ('Научная новизна.', 'Для заданной нелинейной модели с неизвестными характеристиками помощи построен протокол идентификации девяти параметров по девяти точным вероятностям успеха и доказана однозначность восстановления на образе модели.'),
 ('Результаты.', 'Установлены невырожденность якобиана и локальная устойчивость обратного отображения. Восемьдесят восстановлений на синтетических данных и 180000 повторений калибровки показывают влияние коротких интервалов наблюдения и почти постоянной предпосылки на обусловленность.'),
 ('Практическая значимость.', 'Результат задаёт проверяемые требования к диагностическому протоколу и исследовательскому модулю интеллектуальной обучающей системы. Применимость к реальным учебным событиям требует отдельной проверки; точные вероятности не отождествляются с единичными ответами студентов.'),
]
ABSTRACT_EN = [
 ('Relevance.', 'Success on assisted tasks may overestimate independent preparation, making it necessary to distinguish state dynamics from the effect of assistance.'),
 ('Aim.', 'To establish conditions for uniquely recovering the parameters of a two-component model of learning, forgetting and directed interdisciplinary transfer.'),
 ('Methods.', 'The study uses ordinary differential equations, constructive inversion of a finite observation map, sensitivity equations, singular-value analysis and binomial simulation.'),
 ('Scientific novelty.', 'For the specified nonlinear model with unknown assistance characteristics, a protocol identifies nine parameters from nine exact success probabilities, with uniqueness proved on the model image.'),
 ('Results.', 'The observation Jacobian is nonsingular and its inverse is locally stable. Eighty synthetic recovery cases and 180000 calibration replicates demonstrate the conditioning effects of short observation intervals and nearly constant prerequisite states.'),
 ('Practical significance.', 'The results provide verifiable requirements for assessment protocols and a research module of an intelligent learning system. Application to actual learning events requires separate validation; exact probabilities do not represent individual student answers.'),
]
KEYS = 'структурная идентифицируемость, динамическая модель, междисциплинарное обучение, внешняя поддержка, диагностика, численный эксперимент, интеллектуальная обучающая система'
KEYS_EN = 'structural identifiability, dynamical model, interdisciplinary learning, external assistance, assessment, numerical experiment, intelligent learning system'

# Authors are (surname, initials). Metadata were checked against publishers.
REFS = [
 (1, [('Corbett','A.T.'),('Anderson','J.R.')], 'Knowledge tracing: Modeling the acquisition of procedural knowledge', 'User Modeling and User-Adapted Interaction',1994,'4','', '253–278','10.1007/BF01099821'),
 (10,[('Abdelrahman','G.'),('Wang','Q.'),('Nunes','B.')], 'Knowledge tracing: A survey','ACM Computing Surveys',2023,'55','11','Art. 224. P. 1–37','10.1145/3569576'),
 (11,[('Bai','Y.'),('Zhao','J.'),('Wei','T.'),('Cai','Q.'),('He','L.')], 'A survey of explainable knowledge tracing','Applied Intelligence',2024,'54','8','6483–6514','10.1007/s10489-024-05509-8'),
 (12,[('Cui','C.'),('Ma','H.'),('Dong','X.'),('Zhang','C.'),('Zhang','C.'),('Yao','Y.'),('Chen','M.'),('Ma','Y.')], 'Model-agnostic counterfactual reasoning for identifying and mitigating answer bias in knowledge tracing','Neural Networks',2024,'178','','Art. 106495','10.1016/j.neunet.2024.106495'),
 (2,[('Bellman','R.'),('Åström','K.J.')], 'On structural identifiability','Mathematical Biosciences',1970,'7','3–4','329–339','10.1016/0025-5564(70)90132-X'),
 (3,[('Miao','H.'),('Xia','X.'),('Perelson','A.S.'),('Wu','H.')], 'On identifiability of nonlinear ODE models and applications in viral dynamics','SIAM Review',2011,'53','1','3–39','10.1137/090757009'),
 (4,[('Simpson','M.J.'),('Baker','R.E.')], 'Parameter identifiability, parameter estimation, and model prediction for differential equation models','SIAM Review',2026,'68','1','153–171','10.1137/24M1667968'),
 (13,[('Villaverde','A.F.'),('Pathirana','D.'),('Fröhlich','F.'),('Hasenauer','J.'),('Banga','J.R.')], 'A protocol for dynamic model calibration','Briefings in Bioinformatics',2022,'23','1','Art. bbab387','10.1093/bib/bbab387'),
 (14,[('Díaz-Seoane','S.'),('Rey Barreiro','X.'),('Villaverde','A.F.')], 'STRIKE-GOLDD 4.0: user-friendly, efficient analysis of structural identifiability and observability','Bioinformatics',2023,'39','1','Art. btac748','10.1093/bioinformatics/btac748'),
 (15,[('Massonis','G.'),('Banga','J.R.'),('Villaverde','A.F.')], 'Structural identifiability and observability of compartmental models of the COVID-19 pandemic','Annual Reviews in Control',2021,'51','','441–459','10.1016/j.arcontrol.2020.12.001'),
 (5,[('Wang','Y.'),('Huang','W.'),('Gong','M.'),('Geng','X.'),('Liu','T.'),('Zhang','K.'),('Tao','D.')], 'Identifiability and asymptotics in learning homogeneous linear ODE systems from discrete observations','Journal of Machine Learning Research',2024,'25','154','1–50',''),
 (6,[('Kostanian','A.'),('Beklaryan','A.')], 'Identifiability of a dissipative knowledge-dynamics model: exact recovery under designed excitation, degeneration on observational data','arXiv',2026,'','','arXiv:2610.09889v3. Препринт. Версия 3 от 15.10.2026','10.48550/arXiv.2610.09889'),
 (7,[('Potts','D.'),('Tasche','M.')], 'Parameter estimation for exponential sums by approximate Prony method','Signal Processing',2010,'90','5','1631–1642','10.1016/j.sigpro.2009.11.012'),
 (8,[('Hoeffding','W.')], 'Probability inequalities for sums of bounded random variables','Journal of the American Statistical Association',1963,'58','301','13–30','10.1080/01621459.1963.10500830'),
 (9,[('Virtanen','P.'),('Gommers','R.'),('Oliphant','T.E.')], 'SciPy 1.0: fundamental algorithms for scientific computing in Python','Nature Methods',2020,'17','3','261–272','10.1038/s41592-019-0686-2'),
]

def reference(row, english=False):
    old, authors, title, journal, year, vol, issue, pages, doi = row
    names = ', '.join(f'{s} {i}' for s,i in authors)
    if old == 9: names += ' et al.'
    if english:
        if old == 6:
            tail='Preprint, arXiv:2610.09889v3, version 3, 15 October 2026.'
        else:
            tail=f'{journal}, {year}, vol. {vol}'
            if issue: tail+=f', no. {issue}'
            tail+=', '+(pages.replace('Art. ','art. ').replace('. P. ', ', pp. ') if pages.startswith('Art.') else 'pp. '+pages)+'.'
        text=f'{names} {title}. {tail}'
    else:
        if len(authors)<=3 and old!=9:
            lead=f'{authors[0][0]} {authors[0][1]} {title}'
            if len(authors)>1: lead+=' / '+', '.join(f'{i} {s}' for s,i in authors)
        else:
            lead=title+' / '+', '.join(f'{i} {s}' for s,i in authors[:3])+' et al.'
        if old==6: tail=f'Препринт. arXiv:2610.09889v3. Версия 3 от 15.10.2026.'
        else:
            tail=f'{journal}. – {year}. – Vol. {vol}.'
            if issue: tail+=f' – No. {issue}.'
            tail+=' – '+(pages if pages.startswith('Art.') else 'P. '+pages)+'.'
        text=lead+' // '+tail
    if doi:text+=' DOI: '+doi+'.'
    else:text+=' URL: https://jmlr.org/papers/v25/22-1159.html'+(' (accessed: 23 October 2026).' if english else ' (дата обращения: 23.10.2026).')
    return text

def content():
    m=json.loads((B/'manuscript_itvs.json').read_text(encoding='utf-8'))
    blocks=deepcopy(m['blocks'])
    # New sources support specific methodological statements, not calculations.
    blocks[2][1] = ('Различение скрытого состояния и наблюдаемого результата используется в моделях knowledge tracing [1]. Современная классификация этих моделей приведена в обзоре [10], а проблемы интерпретируемости – в [11]. Исследование смещения ответов [12] показывает, что высокая точность прогноза сама по себе не подтверждает восстановление состояния знаний. Структурная идентифицируемость отвечает на вопрос об однозначности восстановления при точных наблюдениях [2]. Её необходимо отличать от практической точности оценивания по конечным шумным данным [3]. Современное обсуждение связи идентифицируемости, оценивания и прогнозирования дано в работе [4].')
    blocks.insert(3,['p', 'Протокол калибровки динамической модели должен согласовывать структуру, доступные наблюдения и процедуру оценивания [13]. Средства анализа структурной идентифицируемости и наблюдаемости представлены, в частности, в STRIKE-GOLDD 4.0 [14]. Анализ семейств моделей ОДУ [15] показывает зависимость восстанавливаемых характеристик от наблюдаемых выходов. Эти работы используются как методологическая основа; биологические интерпретации их параметров на учебный процесс не переносятся.'])
    at=next(i for i,(k,v) in enumerate(blocks) if k=='p' and v.startswith('Для линейных систем'))
    blocks.insert(at+1,['p','В модели [6] межкомпонентный перенос задаётся интенсивностями практики; при фиксированном входе уравнения состояний разделяются. В настоящей постановке скорость освоения второго компонента зависит от текущего состояния первого, поэтому сохраняется динамическая связь состояний. Структурная теорема в [6] опирается на идеализацию наблюдаемой траектории состояния. Здесь рассматриваются конечные наблюдения вероятностей с неизвестными коэффициентами помощи. Таким образом, различаются и механизм переноса, и наблюдаемые величины; доказанный результат ограничен указанной моделью и протоколом.'])
    numbered=m['numbered_equations']; numbered={int(k):v for k,v in numbered.items()}
    mapping={row[0]:i for i,row in enumerate(REFS,1)}
    result=[]
    for k,v in blocks:
        if k in ['p','s']:
            if v.startswith('Вектор неизвестных параметров обозначен ψ;'):continue
            v=re.sub(r'\[(\d+)\]',lambda mt:'['+str(mapping[int(mt[1])])+']',v)
            v=v.replace('условия протокола раздела 2','условия описанного протокола')
            v=v.replace('После упорядочения параметров, указанного в разделе 1,','При указанном порядке параметров')
            v=re.sub(r'(?<=\d)\.(?=\d)',',',v)
            v=v.replace('STRIKE-GOLDD 4,0','STRIKE-GOLDD 4.0')
            v=v.replace('4,647e-13','4,647·10⁻¹³').replace('7,403e-13','7,403·10⁻¹³')
        elif k=='h':v=re.sub(r'^\d+\s+','',v)
        elif k=='e':v=v.replace('^2q(1-κ)','^2κ(1-κ)')
        elif k=='t':
            v['caption']=re.sub(r'^(Таблица \d+) ',r'\1. ',v['caption'])
            for row in v['rows']:
                for i,t in enumerate(row):
                    t=re.sub(r'(?<=\d)\.(?=\d)',',',t)
                    t=re.sub(r'e-(\d+)',lambda mt:'·10⁻'+mt[1].lstrip('0').translate(str.maketrans('0123456789','⁰¹²³⁴⁵⁶⁷⁸⁹')),t)
                    row[i]=t
        result.append([k,v])
    # Place the figure next to the first conditioning table; a later placement
    # would leave a large blank area at the end of the preceding column.
    intro=next(v for k,v in result if k=='p' and v.startswith('На рисунке 1 приведены'))
    fig=next(v for k,v in result if k=='f')
    result=[(k,v) for k,v in result if not (k=='f' or (k=='p' and v==intro))]
    at=next(i for i,(k,v) in enumerate(result) if k=='t' and v['caption'].startswith('Таблица 2.'))
    result[at+1:at+1]=[('p',intro),('f',fig)]
    return dict(title=TITLE,title_en=TITLE_EN,abstract=ABSTRACT,abstract_en=ABSTRACT_EN,
                keywords=KEYS,keywords_en=KEYS_EN,blocks=result,numbered_equations=numbered,
                references=[reference(r) for r in REFS],references_en=[reference(r,True) for r in REFS],
                reference_mapping=mapping,journal='Доклады ТУСУР',specialty='1.2.2 (технические науки)',
                recent_references_2022_2026=sum(r[4]>=2022 for r in REFS))

def make_figure():
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    data=json.loads((B/'results.json').read_text(encoding='utf-8'))
    plt.rcParams.update({'font.family':'Times New Roman','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
    fig,axes=plt.subplots(2,1,figsize=(3.02,4.65),layout='constrained')
    for mode,style in [('M','o-'),('A','s--')]:
        rows=[r for r in data['conditioning'] if r['mode']==mode]
        axes[0].loglog([r['h'] for r in rows],[r['rest_sigma_min'] for r in rows],style,color='black',markerfacecolor='white' if mode=='A' else 'black',markersize=3,linewidth=.85,label=mode)
    axes[0].set_xlabel('Интервал h');axes[0].set_ylabel('Сингулярное число');axes[0].set_title('а  Калибровка первого компонента',fontsize=9);axes[0].legend(frameon=False,fontsize=9)
    rows=data['transfer_excitation']
    axes[1].loglog([r['lambda1'] for r in rows],[r['sigma_min'] for r in rows],'o-',color='black',markersize=3,linewidth=.85)
    axes[1].set_xlabel('Коэффициент забывания λ₁');axes[1].set_ylabel('Сингулярное число');axes[1].set_title('б  Разделение освоения и переноса',fontsize=9)
    for ax in axes:ax.grid(True,which='major',color='0.65',linewidth=.4)
    fig.savefig(B/'conditioning_tusur.png',dpi=350);fig.savefig(B/'conditioning_tusur.svg');plt.close(fig)
    svg=B/'conditioning_tusur.svg'
    svg.write_bytes(('\n'.join(line.rstrip() for line in svg.read_text(encoding='utf-8').splitlines())+'\n').encode('utf-8'))

# Display-line splits preserve the equations and fit the publisher's 8-cm column.
WRAPPED={
 1:r'\begin{aligned}\dot x_1&=\alpha_1u_1(t)(1-x_1)-\lambda_1x_1,\\x_1(0)&=X_1,\end{aligned}',
 2:r'\begin{aligned}\dot x_2&=u_2(t)(\alpha_2+\beta x_1)(1-x_2)-\lambda_2x_2,\\x_2(0)&=X_2.\end{aligned}',
 3:r'\begin{aligned}p_i(t,χ)&=χ\eta_i+(1-χ\eta_i)x_i(t),\\χ&\in\{0,1\},\quad i=1,2.\end{aligned}',
 4:r'\begin{aligned}\psi=&(X_1,\lambda_1,\eta_1,X_2,\lambda_2,\eta_2,\\&\alpha_1,\alpha_2,\beta)\in\Theta,\end{aligned}',
 5:r'\begin{aligned}\Theta=&(0,1)\times(0,\infty)\times(0,1)\\&\times(0,1)\times(0,\infty)\times(0,1)\\&\times(0,\infty)^3.\end{aligned}',
 6:r'\begin{aligned}X&=P_0,\quad κ=P_1/P_0,\\\lambda&=-h^{-1}\ln κ,\end{aligned}',
 8:r'κ=\frac{P_1-P_2}{P_0-P_1},\quad\lambda=-h^{-1}\ln κ,',
 10:r'\begin{aligned}\det J_M&=-hXκ(1-Xκ^2),\\\det J_A&=-hX(1-\eta)^2κ(1-κ)^2.\end{aligned}',
 11:r'\begin{aligned}A=&\frac{\alpha_1c_1}{k}\\&+\left(X_1e^{-2h\lambda_1}-\frac{\alpha_1c_1}{k}\right)e^{-kL},\\k=&\alpha_1c_1+\lambda_1.\end{aligned}',
 12:r'\begin{aligned}w(t)&=Ae^{-\lambda_1t},\\B&=X_2e^{-\lambda_2(2h+L)},\end{aligned}',
 13:r'\begin{aligned}\dot z&=c_2(\alpha+\beta w(t))(1-z)-\lambda_2z,\\z(0)&=B,\quad\alpha=\alpha_2.\end{aligned}',
 14:r'\begin{aligned}\dot v+k(t)v&=c_2(\delta\alpha+\delta\beta w(t))\\&\quad\times(1-\widetilde z(t)),\quad v(0)=0,\end{aligned}',
 16:r'\begin{aligned}S_\alpha(t)&=c_2e^{-K(t)}\int_0^tW(s)\,ds,\\S_\beta(t)&=c_2e^{-K(t)}\int_0^tW(s)w(s)\,ds.\end{aligned}',
 18:r'\begin{aligned}\det J_{\alpha,\beta}=&S_\alpha(t_1)S_\alpha(t_2)\\&\times\bigl(R(t_2)-R(t_1)\bigr)<0.\end{aligned}',
 19:r'\begin{aligned}\det DF=&\det J_1\det J_2(1-\eta_1)\frac{\partial A}{\partial\alpha_1}\\&\times(1-\eta_2)^2\det J_{\alpha,\beta}\ne0.\end{aligned}',
 20:r'\begin{aligned}\|\psi-\widetilde\psi\|_2&\le\frac{2}{\sigma_0}\|F(\psi)-F(\widetilde\psi)\|_2,\\\sigma_0&=\sigma_{\min}(DF(\psi_0))>0.\end{aligned}',
 22:r'\begin{aligned}\|\widehat p-p\|_\infty&\le\sqrt{\frac{\ln(18/\zeta)}{2n}},\\\|\widehat p-p\|_2&\le3\sqrt{\frac{\ln(18/\zeta)}{2n}}.\end{aligned}',
}

def build(template, output, xsl_path):
    m=content();make_figure();doc=Document(template)
    source_paras=[deepcopy(p._p) for p in doc.paragraphs]
    source_sections=[deepcopy(s._sectPr) for s in doc.sections]
    body=doc._element.body
    for child in list(body):body.remove(child)
    def par(text='',slot=14,size=None,bold=None,italic=None,indent=None,align=None):
        node=deepcopy(source_paras[slot]);
        for child in list(node):
            if child.tag!=qn('w:pPr'):node.remove(child)
        for old in node.findall('.//'+qn('w:sectPr')):old.getparent().remove(old)
        body.append(node)
        from docx.text.paragraph import Paragraph
        p=Paragraph(node,doc._body);p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
        p.paragraph_format.widow_control=True
        if indent is not None:p.paragraph_format.first_line_indent=Cm(indent)
        if align is not None:p.alignment=align
        # Italicize inline Latin mathematical symbols; prose and code names stay upright.
        pattern=r'J_[MA]|S[αβ]|Jα,β|ψ|(?<![A-Za-z])(?:[xXuPpTAWBFRSLnhzvgkw][₀₁₂₃₄₅₆₇₈₉]*(?:[αβ])?)(?![A-Za-z])'
        pos=0
        for match in re.finditer(pattern,text) if slot==14 else []:
            if match.start()>pos:p.add_run(text[pos:match.start()])
            token=match[0]
            if token.startswith('J_') or token=='Jα,β':
                run=p.add_run('J');run.bold=True;run.italic=False
                run=p.add_run(token[2:] if token.startswith('J_') else 'α,β');run.font.subscript=True;run.bold=True;run.italic=False
            elif token in ['Sα','Sβ']:
                run=p.add_run('S');run.italic=True
                run=p.add_run(token[1]);run.font.subscript=True;run.italic=False
            elif token=='ψ':run=p.add_run(token);run.bold=True;run.italic=False
            else:run=p.add_run(token);run.italic=True
            pos=match.end()
        p.add_run(text[pos:])
        for run in p.runs:
            run.font.color.rgb=RGBColor(0,0,0)
            if slot not in [4,6]:run.font.name='Times New Roman'
            if size:run.font.size=Pt(size)
            if bold is not None:run.bold=bold
            if italic is not None:run.italic=italic
        return p
    def meta(items,slot):
        p=par('',slot=slot,size=9,indent=0)
        for i,(label,text) in enumerate(items):
            if i:p.add_run(' ')
            r=p.add_run(label+' ');r.bold=True;r.italic=True
            p.add_run(text)
        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9)
        return p
    par('УДК 519.876.5',1,size=10,indent=0)
    par('Шифр специальности ВАК 1.2.2 (технические науки)',2,size=10,indent=0)
    par('',slot=5,size=4)
    par('П.В. Никитин',4,indent=0)
    par('',slot=5,size=4)
    par(TITLE,6,indent=0)
    par('',slot=7,size=4)
    meta(ABSTRACT,8)
    p=par('Ключевые слова: '+KEYS+'.',9,size=9,indent=0)
    for r in p.runs:r.bold=False
    p.paragraph_format.keep_with_next=True
    # Continuous section transition preserves the publisher's one/two-column geometry.
    boundary=par('',slot=12,size=1,indent=0)
    s=deepcopy(source_sections[0]);
    for n in s.findall(qn('w:pgNumType')):n.set(qn('w:start'),'1')
    boundary._p.get_or_add_pPr().append(s)
    transform=E.XSLT(E.parse(str(xsl_path)));formula_index=0
    exports=[]
    for kind,value in m['blocks']:
        if kind=='h':
            p=par(value,13,size=10,bold=True);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
        elif kind in ['p','s']:
            p=par(value,14,size=10)
            if kind=='s':
                # Only the theorem/lemma label is bold, keeping the proof compact.
                txt=p.text;p.clear();mt=re.match(r'^((?:Лемма|Теорема|Предложение) \d+\.)',txt)
                if mt:r=p.add_run(mt[0]);r.bold=True;p.add_run(txt[len(mt[0]):])
                else:p.add_run(txt)
                for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10)
                p.paragraph_format.keep_with_next=True
        elif kind=='e':
            formula_index+=1;tex=WRAPPED.get(formula_index,value)
            if tex.startswith(r'\begin{aligned}'):
                ns='http://www.w3.org/1998/Math/MathML'
                root=E.Element('{'+ns+'}math',nsmap={None:ns})
                table=E.SubElement(root,'{'+ns+'}mtable',columnalign='left')
                inner=tex[len(r'\begin{aligned}'):-len(r'\end{aligned}')]
                for line in re.split(r'\\\\',inner):
                    cell=E.SubElement(E.SubElement(table,'{'+ns+'}mtr'),'{'+ns+'}mtd')
                    parsed=E.fromstring(convert(line.replace('&','')).encode())
                    for child in parsed:cell.append(deepcopy(child))
                mathml=E.tostring(root,encoding='unicode')
            else:mathml=convert(tex)
            math=transform(E.fromstring(mathml.encode())).getroot()
            if math.tag==qn('m:oMathPara'):math=math.find(qn('m:oMath'))
            for run in math.findall('.//'+qn('m:r')):
                rp=run.find(qn('w:rPr'))
                if rp is None:rp=OxmlElement('w:rPr');run.insert(0,rp)
                ff=OxmlElement('w:rFonts');ff.set(qn('w:ascii'),'Times New Roman');ff.set(qn('w:hAnsi'),'Times New Roman');rp.append(ff)
                sz=OxmlElement('w:sz');sz.set(qn('w:val'),'20');rp.append(sz)
                text=''.join(run.itertext())
                if re.search('[α-ωΑ-ΩχκψΘ]',text):
                    mrp=run.find(qn('m:rPr'))
                    if mrp is None:mrp=OxmlElement('m:rPr');run.insert(0,mrp)
                    sty=OxmlElement('m:sty');sty.set(qn('m:val'),'b' if 'ψ' in text else 'p');mrp.append(sty)
                if text in ['det','ln','min','exp']:
                    mrp=run.find(qn('m:rPr'))
                    if mrp is None:mrp=OxmlElement('m:rPr');run.insert(0,mrp)
                    sty=OxmlElement('m:sty');sty.set(qn('m:val'),'p');mrp.append(sty)
            p=par('',14,size=10,indent=0,align=WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.keep_together=True;p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
            p._p.append(math)
            if formula_index in m['numbered_equations']:
                p.add_run('\t('+str(m['numbered_equations'][formula_index])+')')
                tabs=OxmlElement('w:tabs');tab=OxmlElement('w:tab');tab.set(qn('w:val'),'right');tab.set(qn('w:pos'),'4535');tabs.append(tab);p._p.get_or_add_pPr().append(tabs)
            exports.append({'index':formula_index,'number':m['numbered_equations'].get(formula_index),'tex':tex,'semantic_tex':value,'mathml':mathml})
        elif kind=='t':
            number=int(re.search(r'\d+',value['caption'])[0]);caption=par(value['caption'],14,size=9,indent=0)
            caption.paragraph_format.keep_with_next=True;caption.paragraph_format.space_before=Pt(4)
            heads=value['heads']
            if number==2:heads=['h','Режим','σmin\nкалибровки','σmin\nполного DF']
            if number==3:heads=['h','Режим','Вне образа, %','Медиана\nошибки η']
            widths=[1.5,1.7,4.8] if number==1 else [0.75,1.15,2.7,3.4]
            tb=doc.add_table(rows=1,cols=len(heads));tb.alignment=WD_TABLE_ALIGNMENT.CENTER;tb.autofit=False;body.append(tb._tbl)
            for c,w in zip(tb.columns,widths):c.width=Cm(w)
            for c,v in zip(tb.rows[0].cells,heads):c.text=v
            for row in value['rows']:
                for c,v in zip(tb.add_row().cells,row):c.text=v
            for ri,row in enumerate(tb.rows):
                trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
                if ri==0:trpr.append(OxmlElement('w:tblHeader'))
                for ci,c in enumerate(row.cells):
                    c.width=Cm(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;pr=c._tc.get_or_add_tcPr()
                    for sh in list(pr.findall(qn('w:shd'))):pr.remove(sh)
                    bd=OxmlElement('w:tcBorders')
                    for edge in ['top','bottom','left','right']:
                        n=OxmlElement('w:'+edge);n.set(qn('w:val'),'single');n.set(qn('w:sz'),'4');n.set(qn('w:color'),'000000');bd.append(n)
                    pr.append(bd)
                    mar=OxmlElement('w:tcMar')
                    for edge,sz in [('top','35'),('bottom','35'),('left','40'),('right','40')]:
                        n=OxmlElement('w:'+edge);n.set(qn('w:w'),sz);n.set(qn('w:type'),'dxa');mar.append(n)
                    pr.append(mar)
                    for p in c.paragraphs:
                        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=1
                        p.alignment=WD_ALIGN_PARAGRAPH.LEFT if number==1 else WD_ALIGN_PARAGRAPH.CENTER
                        for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(9);r.bold=False
            p=par('',14,size=3,indent=0)
        elif kind=='f':
            p=par('',14,size=9,indent=0,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_with_next=True
            p.add_run().add_picture(str(B/'conditioning_tusur.png'),width=Cm(7.85))
            p=par(value['caption'],14,size=9,indent=0,align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.keep_together=True
    p=par('Литература',171,size=10,italic=True);p.paragraph_format.keep_with_next=True
    for i,r in enumerate(m['references'],1):par(str(i)+'. '+r,172,size=9)
    par('',14,size=4)
    par('Никитин Пётр Владимирович',130,size=9,bold=True,indent=0)
    for t in ['Канд. пед. наук, доцент кафедры искусственного интеллекта Финансового университета при Правительстве Российской Федерации',
              'Ленинградский пр-т, 49/2, Москва, Россия, 125167',
              'ORCID: 0000-0001-8866-5610','Тел. организации: +7 (499) 700-86-00','Эл. почта: pvnikitin@fa.ru']:
        par(t,131,size=9,indent=0)
    par('',14,size=5)
    par('Nikitin P.V.',164,size=9,indent=0)
    par(TITLE_EN,165,size=9,bold=True,indent=0)
    meta(ABSTRACT_EN,167)
    par('Keywords: '+KEYS_EN+'.',168,size=9,bold=False,indent=0)
    p=par('References',171,size=10,italic=True);p.paragraph_format.keep_with_next=True
    for i,r in enumerate(m['references_en'],1):par(str(i)+'. '+r,172,size=9)
    par('',14,size=4)
    par('Petr V. Nikitin',190,size=9,bold=True,indent=0)
    for t in ['Candidate of Pedagogical Sciences, Associate Professor, Department of Artificial Intelligence, Financial University under the Government of the Russian Federation',
              '49/2, Leningradsky Ave., Moscow, Russia, 125167',
              'ORCID: 0000-0001-8866-5610','Institution phone: +7 (499) 700-86-00','Email: pvnikitin@fa.ru']:
        par(t,191,size=9,indent=0)
    final_sec=deepcopy(source_sections[1])
    for node in final_sec.findall(qn('w:pgNumType')):final_sec.remove(node)
    body.append(final_sec)
    doc.core_properties.title=TITLE;doc.core_properties.author='Пётр Владимирович Никитин';doc.core_properties.subject='Авторская редакция для журнала Доклады ТУСУР'
    output.parent.mkdir(parents=True,exist_ok=True)
    # Repack from the retained template so untouched parts remain byte-identical.
    with tempfile.TemporaryDirectory(prefix='tusur-',dir=output.parent) as td:
        temp=Path(td)/'built.docx';doc.save(temp)
        with ZipFile(template) as original,ZipFile(temp) as built:
            edit={'word/document.xml','word/_rels/document.xml.rels','docProps/core.xml','[Content_Types].xml'}
            changed={n:built.read(n) for n in edit}
            for name in original.namelist():
                if re.fullmatch(r'word/(?:header|footer)\d+\.xml',name):
                    root=E.fromstring(original.read(name));paras=root.findall('.//'+qn('w:p'))
                    for pi,p in enumerate(paras):
                        for ch in list(p):
                            if ch.tag!=qn('w:pPr'):p.remove(ch)
                        if pi:continue
                        run=OxmlElement('w:r');rp=OxmlElement('w:rPr');font=OxmlElement('w:rFonts');font.set(qn('w:ascii'),'Times New Roman');font.set(qn('w:hAnsi'),'Times New Roman');rp.append(font);sz=OxmlElement('w:sz');sz.set(qn('w:val'),'18');rp.append(sz);run.append(rp)
                        tt=OxmlElement('w:t');tt.text=('П.В. Никитин. Идентификация динамической модели' if 'header' in name else 'Доклады ТУСУР');run.append(tt);p.append(run)
                        if 'header' in name:
                            tabrun=OxmlElement('w:r');tabrun.append(OxmlElement('w:tab'));p.append(tabrun)
                            fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p.append(fld)
                    changed[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
            with ZipFile(output,'w',ZIP_DEFLATED) as final:
                for info in original.infolist():final.writestr(info,changed.get(info.filename,original.read(info.filename)))
                for name in built.namelist():
                    if name not in original.namelist():final.writestr(name,built.read(name))
    (B/'manuscript_tusur.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')
    (B/'equations_tusur.json').write_text(json.dumps(exports,ensure_ascii=False,indent=2),encoding='utf-8')
    (B/'equations_tusur.tex').write_text('\n\n'.join('% '+str(r['index'])+'\n\\[\n'+r['tex']+'\n\\]' for r in exports),encoding='utf-8')
    mathdir=B/'mathml_tusur';mathdir.mkdir(exist_ok=True)
    for r in exports:(mathdir/f'equation-{r["index"]:02d}.mml').write_text(r['mathml'],encoding='utf-8')
    print(json.dumps({'file':str(output),'formulas':formula_index,'references':len(REFS),'recent_2022_2026':m['recent_references_2022_2026'],'abstract_words_ru':sum(len(t.split()) for _,t in ABSTRACT),'abstract_words_en':sum(len(t.split()) for _,t in ABSTRACT_EN),'math_type_compliant':False},ensure_ascii=False,indent=2))

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--template',type=Path,required=True);parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--xsl',type=Path,default=Path('C:/Program Files/Microsoft Office/root/Office16/MML2OMML.XSL'))
    args=parser.parse_args();build(args.template,args.output,args.xsl)
