from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import json, re

BASE=Path(__file__).parent
import os
OUT=Path(os.environ.get('DOCTOR_PIT_OUTPUT_DIR',str(BASE.parent.parent/'outputs')))
OUT.mkdir(parents=True,exist_ok=True)
FINAL=OUT/'Статья_EJCE_Междисциплинарные_траектории_русская.docx'
DATA=json.loads((BASE/'manuscript.json').read_text(encoding='utf8'))
doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2)
sec.header_distance=sec.footer_distance=Cm(.8)

def font(f,name='Times New Roman',size=14):
    f.name=name;f.size=Pt(size);f.bold=False;f.italic=False
    f.color.rgb=RGBColor(0,0,0)

for name in ['Normal','Title','Heading 1','Heading 2','Caption','Body Text','No Spacing']:
    st=doc.styles[name];font(st.font)
    fmt=st.paragraph_format
    fmt.line_spacing=1.5;fmt.space_before=Pt(0);fmt.space_after=Pt(0)
    fmt.first_line_indent=Cm(1);fmt.widow_control=True
    if name in ['Title','Heading 1','Heading 2']:
        fmt.first_line_indent=Cm(0);fmt.keep_with_next=True
        fmt.space_before=Pt(9);fmt.space_after=Pt(3)
    # Theme font bindings would override an explicit font in some Word versions.
    rp=st.element.get_or_add_rPr()
    for unwanted in ['w:spacing','w:kern']:
        for node in list(rp.findall(qn(unwanted))):rp.remove(node)
    size_cs=rp.find(qn('w:szCs'))
    if size_cs is not None:size_cs.set(qn('w:val'),'28')
    for node in list(st.element.findall('.//'+qn('w:pBdr'))):node.getparent().remove(node)
    rf=rp.find(qn('w:rFonts'))
    if rf is not None:
        for attr in ['asciiTheme','hAnsiTheme','csTheme','eastAsiaTheme']:
            rf.attrib.pop(qn('w:'+attr),None)
        rf.set(qn('w:eastAsia'),'Times New Roman')
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');rp.append(lang)

normal=doc.styles['Normal']
normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
doc.core_properties.title=DATA['title']
doc.core_properties.subject='Педагогическое проектирование междисциплинарных траекторий'
doc.core_properties.author='Никитин Пётр Владимирович'
doc.core_properties.keywords=DATA['keywords']
doc.core_properties.comments=''

def p(text='',style=None,align=None,indent=1,before=0,after=0,keep=False):
    x=doc.add_paragraph(style=style)
    if text:
        r=x.add_run(text);font(r.font)
    f=x.paragraph_format
    f.first_line_indent=Cm(indent);f.space_before=Pt(before);f.space_after=Pt(after)
    f.keep_with_next=keep;f.widow_control=True
    if align is not None:f.alignment=align
    return x

def xml(tag,attrs=None,text=None):
    x=OxmlElement(tag)
    for key,value in (attrs or {}).items():x.set(qn(key),str(value))
    if text is not None:x.text=text
    return x

def rpr(size=14,math=False):
    x=xml('w:rPr')
    x.append(xml('w:rFonts',{'w:ascii':'Cambria Math' if math else 'Times New Roman','w:hAnsi':'Cambria Math' if math else 'Times New Roman','w:cs':'Times New Roman'}))
    x.append(xml('w:sz',{'w:val':size*2}));x.append(xml('w:szCs',{'w:val':size*2}))
    x.append(xml('w:b',{'w:val':'0'}));x.append(xml('w:i',{'w:val':'0'}))
    x.append(xml('w:color',{'w:val':'000000'}));return x

def mr(text):
    x=xml('m:r');mp=xml('m:rPr');mp.append(xml('m:sty',{'m:val':'p'}))
    x.append(mp);x.append(rpr(math=True));x.append(xml('m:t',text=text));return x

def wrap(tag,items):
    x=xml('m:'+tag)
    for item in items:x.append(mr(item) if isinstance(item,str) else item)
    return x

def sub(b,i):return wrap('sSub',[wrap('e',[b]),wrap('sub',[i])])
def sup(b,i):return wrap('sSup',[wrap('e',[b]),wrap('sup',[i])])
def subsup(b,i,j):return wrap('sSubSup',[wrap('e',[b]),wrap('sub',[i]),wrap('sup',[j])])
def frac(a,b):return wrap('f',[wrap('num',a if isinstance(a,list) else [a]),wrap('den',b if isinstance(b,list) else [b])])
def nary(char,lower,upper,items):
    x=xml('m:nary');pr=xml('m:naryPr')
    pr.append(xml('m:chr',{'m:val':char}));pr.append(xml('m:limLoc',{'m:val':'subSup'}));x.append(pr)
    x.append(wrap('sub',[lower]));x.append(wrap('sup',[upper]));x.append(wrap('e',items));return x

def eqline(items,number=None,keep=False):
    x=p('',indent=0,before=2,after=2,keep=keep)
    f=x.paragraph_format;f.line_spacing=1.5
    tabs=f.tab_stops;tabs.add_tab_stop(Cm(8.5),WD_TAB_ALIGNMENT.CENTER);tabs.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT)
    x.add_run('\t')
    om=xml('m:oMath')
    for item in items:om.append(mr(item) if isinstance(item,str) else item)
    x._p.append(om)
    if number is not None:
        x.add_run('\t');font(x.add_run('('+str(number)+')').font)
    return x

def equation(n):
    if n==1:eqline(['k = ',frac('N(1 − r)','d')],n)
    if n==2:
        eqline([sub('X','j'),' ∼ Bin(k, ',sub('θ','j'),'),  ',sub('Y','j'),' ∼ Bin(Nr, γ),  j ∈ {1, 2};'],keep=True)
        eqline(['Z ∼ Bin(5, φ)'],n)
    if n==3:
        eqline([sub('R','0'),': d',sub('X','j'),' + ',sub('Y','j'),' ≥ ⌈0,8N⌉,  j ∈ {1, 2};'],keep=True)
        eqline([sub('R','1'),': ',sub('X','j'),' ≥ ⌈0,8k⌉,  j ∈ {1, 2},  Z ≥ 4;'],keep=True)
        eqline([sub('R','2'),': ',sub('q','j'),' ≥ 0,8,  j ∈ {1, 2},  Z ≥ 4'],n)
    if n==4:
        eqline(['Pr(',sub('R','ℓ'),') = T(φ)',nary('∏','j=1','2',['Pr(',sub('X','j'),' ≥ ',sub('h','ℓ'),')']),',  ℓ ∈ {1, 2};'],keep=True)
        eqline(['T(φ) = ',nary('∑','z=4','5',[subsup('C','5','z'),sup('φ','z'),sup('(1 − φ)','5−z')])],n)

footnote_items=[]
def footnote(par,text):
    i=len(footnote_items)+1;footnote_items.append((i,text))
    r=xml('w:r');pr=rpr();pr.append(xml('w:vertAlign',{'w:val':'superscript'}));r.append(pr)
    r.append(xml('w:footnoteReference',{'w:id':i}));par._p.append(r)

def table(b):
    cap=p('Таблица '+str(b['number'])+'. '+b['title'],align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,before=6,after=3,keep=True)
    t=doc.add_table(rows=1,cols=len(b['headers']))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,b['widths']):c.width=Cm(w)
    pr=t._tbl.tblPr
    borders=xml('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        borders.append(xml('w:'+edge,{'w:val':'single','w:sz':'4','w:color':'808080'}))
    pr.append(borders)
    mar=xml('w:tblCellMar')
    for edge in ['top','bottom','left','right']:mar.append(xml('w:'+edge,{'w:w':'65','w:type':'dxa'}))
    pr.append(mar)
    for ri,values in enumerate([b['headers']]+b['rows']):
        row=t.rows[0] if ri==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr();trpr.append(xml('w:cantSplit'))
        if ri==0:trpr.append(xml('w:tblHeader'))
        for ci,(cell,text,width) in enumerate(zip(row.cells,values,b['widths'])):
            cell.width=Cm(width);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=cell._tc.get_or_add_tcPr()
            if ri==0:cp.append(xml('w:shd',{'w:fill':'F2F2F2'}))
            par=cell.paragraphs[0];par.paragraph_format.first_line_indent=Cm(0)
            par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(0)
            par.paragraph_format.line_spacing=1.5
            par.paragraph_format.keep_with_next=ri<len(b['rows'])
            par.alignment=WD_ALIGN_PARAGRAPH.LEFT if b['number']==1 and ci==4 else WD_ALIGN_PARAGRAPH.CENTER
            font(par.add_run(text).font)

def figure(b):
    par=p('',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,before=6,keep=True)
    par.paragraph_format.line_spacing=1
    shape=par.add_run().add_picture(str(BASE/b['file']),width=Cm(16.8))
    shape._inline.docPr.set('descr',b['title'])
    shape._inline.docPr.set('title','Рисунок '+str(b['number']))
    p('Рис. '+str(b['number'])+'. '+b['title'],align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,before=3,after=6)

p('УДК 378.147:004.8',align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,after=6)
p(DATA['title'],style='Title',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=9,keep=True)
p('Никитин Пётр Владимирович',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=6,keep=True)
p('Финансовый университет при Правительстве Российской Федерации, Москва, Российская Федерация',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,keep=True)
p('Кандидат педагогических наук, доцент; кафедра искусственного интеллекта, факультет информационных технологий и анализа больших данных',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=3)
p('125167, Москва, Ленинградский проспект, д. 49/2',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0)
p('E-mail: pvnikitin@fa.ru; ORCID: 0000-0001-8866-5610',align=WD_ALIGN_PARAGRAPH.CENTER,indent=0,after=9)
p('Аннотация',indent=0,keep=True,after=3)
p(DATA['abstract'])
p('Ключевые слова: '+DATA['keywords']+'.',before=3,after=6)

for b in DATA['blocks']:
    if b['type'] in ['heading','subheading']:
        p(b['text'],style='Heading 1' if b['type']=='heading' else 'Heading 2',align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,before=9,after=3,keep=True)
    elif b['type']=='paragraph':
        par=p(b['text'])
        if 'footnote' in b:footnote(par,b['footnote'])
    elif b['type']=='equation':equation(b['id'])
    elif b['type']=='table':table(b)
    elif b['type']=='figure':figure(b)

p('Literature',style='Heading 1',align=WD_ALIGN_PARAGRAPH.LEFT,indent=0,before=9,after=6,keep=True)
for ref in DATA['references']:
    par=p(ref['text'],indent=0,after=3)
    par.paragraph_format.left_indent=Cm(1)
    par.paragraph_format.first_line_indent=Cm(-1)
    par.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT

OUT.mkdir(exist_ok=True)
doc.save(FINAL)

# Create native editable Word footnotes; python-docx does not expose this part.
with ZipFile(FINAL) as z:parts={n:z.read(n) for n in z.namelist()}
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
REL='http://schemas.openxmlformats.org/package/2006/relationships'
CT='http://schemas.openxmlformats.org/package/2006/content-types'
root=E.Element('{'+W+'}footnotes',nsmap={'w':W})
for i,typ,sep in [(-1,'separator','separator'),(0,'continuationSeparator','continuationSeparator')]:
    fn=xml('w:footnote',{'w:id':i,'w:type':typ});pp=xml('w:p');rr=xml('w:r');rr.append(xml('w:'+sep));pp.append(rr);fn.append(pp);root.append(fn)
for i,text in footnote_items:
    fn=xml('w:footnote',{'w:id':i});pp=xml('w:p');pr=xml('w:pPr')
    pr.append(xml('w:spacing',{'w:before':'0','w:after':'0','w:line':'360','w:lineRule':'auto'}));pp.append(pr)
    rr=xml('w:r');rp=rpr();rp.append(xml('w:vertAlign',{'w:val':'superscript'}));rr.append(rp);rr.append(xml('w:footnoteRef'));pp.append(rr)
    rr=xml('w:r');rr.append(rpr());tt=xml('w:t',text=' '+text);tt.set('{http://www.w3.org/XML/1998/namespace}space','preserve');rr.append(tt);pp.append(rr)
    fn.append(pp);root.append(fn)
parts['word/footnotes.xml']=E.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)
rels=E.fromstring(parts['word/_rels/document.xml.rels'])
rid='rIdFootnotesArticle'
rels.append(E.Element('{'+REL+'}Relationship',Id=rid,Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',Target='footnotes.xml'))
parts['word/_rels/document.xml.rels']=E.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
ct=E.fromstring(parts['[Content_Types].xml'])
ct.append(E.Element('{'+CT+'}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml'))
parts['[Content_Types].xml']=E.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
    for n,content in parts.items():z.writestr(n,content)
print(FINAL)
print('Tables',len(doc.tables),'references',len(DATA['references']),'native footnotes',len(footnote_items))
