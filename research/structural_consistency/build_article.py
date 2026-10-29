"""Build the manuscript from its source and frozen experimental results.

Requirements: python-docx, lxml, matplotlib; Word's MML2OMML.XSL on Windows.
The scientific experiment itself has no third-party dependencies.
"""
from pathlib import Path
import sys,os,json,re,statistics,math,copy
B=Path(__file__).resolve().parent
if Path(r'D:\CodexWork\2026-10-22\new-chat\diff-eq-deps').exists():
 sys.path.insert(0,r'D:\CodexWork\2026-10-22\new-chat\diff-eq-deps')
os.environ.setdefault('MPLCONFIGDIR',str(B/'mpl-cache'))
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle,FancyArrowPatch
from lxml import etree
from xml.sax.saxutils import escape
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_BREAK,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT=Path(sys.argv[1]) if len(sys.argv)>1 else B/'deliverables'
OUT.mkdir(parents=True,exist_ok=True)
NAME='Никитин_Структурная_согласованность_Вестник_НГУ_2026-10-25.docx'
R=json.loads((B/'results.json').read_text(encoding='utf8'))
M=json.loads((B/'source_manifest.json').read_text(encoding='utf8'))

refs=[
('Dawson S., Hubball H. Curriculum analytics: Application of social network analysis for improving strategic curriculum decision-making in a research-intensive university. Teaching & Learning Inquiry, 2014, vol. 2, no. 2, pp. 59–74.','10.20343/teachlearninqu.2.2.59'),
('Stavrinides P., Zuev K. M. Course-prerequisite networks for analyzing and understanding academic curricula. Applied Network Science, 2023, vol. 8, article 19.','10.1007/s41109-023-00543-w'),
('Yang B., Gharebhaygloo M., Rondi H. R., Hortis E., Zeledon Lostalo E., Huang X., Ercal G. Comparative analysis of course prerequisite networks for five Midwestern public institutions. Applied Network Science, 2024, vol. 9, article 25.','10.1007/s41109-024-00637-z'),
('Hou B., Lin Y., Li Y., Fang C., Li C., Wang X. KG-PLPPM: A Knowledge Graph-Based Personal Learning Path Planning Method Used in Online Learning. Electronics, 2025, vol. 14, no. 2, article 255.','10.3390/electronics14020255'),
('Xu Z., Guan X., Shi C., Chen Q., Yu R. Evaluating 21st-Century Competencies in Postsecondary Curricula with Large Language Models: Performance Benchmarking and Reasoning-Based Prompting Strategies. Journal of Learning Analytics, 2026, vol. 13, no. 1, pp. 7–30.','10.18608/jla.2026.9127'),
('Winkler S., von Pilgrim J. A survey of traceability in requirements engineering and model-driven development. Software and Systems Modeling, 2010, vol. 9, no. 4, pp. 529–565.','10.1007/s10270-009-0145-0'),
('Hogan A., Blomqvist E., Cochez M. et al. Knowledge Graphs. ACM Computing Surveys, 2021, vol. 54, no. 4, article 71, pp. 1–37.','10.1145/3447772'),
('Gallo G., Longo G., Pallottino S., Nguyen S. Directed hypergraphs and applications. Discrete Applied Mathematics, 1993, vol. 42, no. 2–3, pp. 177–201.','10.1016/0166-218X(93)90045-P'),
('Dowling W. F., Gallier J. H. Linear-time algorithms for testing the satisfiability of propositional Horn formulae. Journal of Logic Programming, 1984, vol. 1, no. 3, pp. 267–284.','10.1016/0743-1066(84)90014-1'),
('Imieliński T., Lipski W. Incomplete Information in Relational Databases. Journal of the ACM, 1984, vol. 31, no. 4, pp. 761–791.','10.1145/1634.1886'),
('Gupta A., Mumick I. S., Subrahmanian V. S. Maintaining views incrementally. Proceedings of the 1993 ACM SIGMOD International Conference on Management of Data. New York, ACM, 1993, pp. 157–166.','10.1145/170035.170066'),
('Motik B., Nenov Y., Piro R., Horrocks I. Incremental Update of Datalog Materialisation: the Backward/Forward Algorithm. Proceedings of the AAAI Conference on Artificial Intelligence, 2015, vol. 29, no. 1, pp. 1560–1568.','10.1609/aaai.v29i1.9409'),
('Corman J., Florenzano F., Reutter J. L., Savković O. Validating SHACL Constraints over a SPARQL Endpoint. The Semantic Web – ISWC 2019. Lecture Notes in Computer Science, vol. 11778. Cham, Springer, 2019, pp. 145–163.','10.1007/978-3-030-30793-6_9'),
('Knublauch H., Kontokostas D. (eds.). Shapes Constraint Language (SHACL). W3C Recommendation, 20 July 2017. URL: https://www.w3.org/TR/shacl/ (accessed 25.10.2026).',None),
('Moreau L., Missier P. (eds.). PROV-DM: The PROV Data Model. W3C Recommendation, 30 April 2013. URL: https://www.w3.org/TR/prov-dm/ (accessed 25.10.2026).',None),
('Chen T. Y., Kuo F.-C., Liu H., Poon P.-L., Towey D., Tse T. H., Zhou Z. Q. Metamorphic Testing: A Review of Challenges and Opportunities. ACM Computing Surveys, 2018, vol. 51, no. 1, article 4, pp. 1–27.','10.1145/3143561')]

def figure():
 plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'svg.fonttype':'none','pdf.fonttype':42,'ps.fonttype':42})
 fig,ax=plt.subplots(figsize=(6.1,2.12));ax.set_xlim(0,10);ax.set_ylim(0,4);ax.axis('off')
 boxes=[(.05,2.35,2.6,1.3,'Учебные документы\nСроки и версии\nИсточники связей'),(3.35,2.35,2.7,1.3,'Утверждённые E⁻\nКандидатные E⁺\nE⁻ ⊆ E⁺'),(6.8,2.35,3.1,1.3,'Два замыкания\nC⁻ ⊆ C⁺\nКонечные выводы'),(3.35,.1,2.7,1.3,'Экспертное уточнение\nсвязей'),(6.8,.1,3.1,1.3,'Обеспечено\nНе обеспечено в модели\nНедостаточно сведений')]
 for x,y,w,h,label in boxes:
  ax.add_patch(Rectangle((x,y),w,h,fill=False,lw=.8));ax.text(x+w/2,y+h/2,label,ha='center',va='center',linespacing=1.35)
 for a,b in [((2.65,3),(3.35,3)),((6.05,3),(6.8,3)),((8.35,2.35),(8.35,1.4)),((6.8,.75),(6.05,.75)),((4.7,1.4),(4.7,2.35))]:
  ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=10,lw=.8,color='black'))
 fig.subplots_adjust(left=.015,right=.99,top=.99,bottom=.025)
 for ext in ['png','eps','svg']:fig.savefig(B/f'figure_1_pipeline.{ext}',dpi=260,facecolor='white')
 plt.close(fig)
figure()

def mi(x):
 # Journal requires upright Greek and italic Latin; MathML handles Latin by default.
 attr=' mathvariant="normal"' if any('\u0370'<=c<='\u03ff' for c in x) else ''
 return f'<mi{attr}>{escape(x)}</mi>'
def mo(x):return '<mo>'+escape(x)+'</mo>'
def tx(x):return '<mtext>'+escape(x)+'</mtext>'
def mn(x):return '<mn>'+str(x)+'</mn>'
def row(*x):return '<mrow>'+''.join(x)+'</mrow>'
def sub(a,b):return '<msub>'+a+b+'</msub>'
def sup(a,b):return '<msup>'+a+b+'</msup>'
def par(a):return row(mo('('),a,mo(')'))
def brace(a):return row(mo('{'),a,mo('}'))
def table(rows):return '<mtable columnalign="left">'+''.join('<mtr>'+''.join('<mtd>'+c+'</mtd>' for c in r)+'</mtr>' for r in rows)+'</mtable>'
def comma(*x):return row(*[s for i,a in enumerate(x) for s in ((mo(','),a) if i else (a,))])
C=mi('C');E=mi('E');V=mi('V');RR=mi('R');r=mi('r');e=mi('e')
Cminus=sup(C,mo('−'));Cp=sup(C,mo('+'));Em=sup(E,mo('−'));Ep=sup(E,mo('+'));C0=sub(C,mn(0));Cr=sub(C,r);Ie=sub(mi('I'),e);oe=sub(mi('o'),e)
cl=lambda ee:sub(tx('cl'),ee)+par(C0)
unionr='<munder>'+mo('⋃')+row(r,mo('≥'),mn(0))+'</munder>'+Cr
eqs={
1:row(mi('ξ'),mo('='),par(comma(mi('k'),mi('ℓ'),mi('σ'),mi('t'),mi('v')))),
2:row(Em,mo('⊆'),E,mo('⊆'),Ep),
3:table([[row(sub(C,row(r,mo('+'),mn(1))),mo('='),Cr,mo('∪'),brace(row(oe,mo(':'),e,mo('∈'),E,mo(','),Ie,mo('⊆'),Cr)))],[row(cl(E),mo('='),unionr)]]),
4:row(Cminus,mo('='),cl(Em),mo('⊆'),cl(E),mo('⊆'),cl(Ep),mo('='),Cp),
5:row(mi('S'),par(RR),mo('='),'<mfenced open="{" close="">'+table([[tx('ensured,'),row(RR,mo('⊆'),Cminus)],[tx('not_ensured,'),row(RR,mo('⊈'),Cp)],[tx('undetermined,'),tx('в остальных случаях')]])+'</mfenced>'),
6:table([[row(mi('B'),mo('='),V,mo('∖'),Cp,mo(','),mi('B'),mo('∩'),C0,mo('='),mo('∅'),mo(','),mi('B'),mo('∩'),RR,mo('≠'),mo('∅'))],[row(mo('∀'),e,mo('∈'),Ep,mo(':'),oe,mo('∈'),mi('B'),mo('⇒'),Ie,mo('∩'),mi('B'),mo('≠'),mo('∅'))]]),
7:row(mi('L'),mo('='),'<munder>'+mo('∑')+row(e,mo('∈'),E)+'</munder>',mo('|'),Ie,mo('|'),mo(','),mi('T'),mo('='),tx('O'),par(row(mo('|'),V,mo('|'),mo('+'),mo('|'),E,mo('|'),mo('+'),mi('L')))),
8:table([[row(sub(mi('D'),row(r,mo('+'),mn(1))),mo('='),sub(mi('D'),r),mo('∪'),brace(row(oe,mo(':'),e,mo('∈'),mi('F'),mo(','),Ie,mo('∩'),sub(mi('D'),r),mo('≠'),mo('∅'))))],[row(mi('D'),mo('='),'<munder>'+mo('⋃')+row(r,mo('≥'),mn(0))+'</munder>',sub(mi('D'),r))]])}
eq_plain={1:'ξ = (k, ℓ, σ, t, v)',2:'E⁻ ⊆ E ⊆ E⁺',3:'Cᵣ₊₁ = Cᵣ ∪ {oₑ : e ∈ E, Iₑ ⊆ Cᵣ}; cl_E(C₀) = ⋃ᵣ≥₀ Cᵣ',4:'C⁻ = cl_E⁻(C₀) ⊆ cl_E(C₀) ⊆ cl_E⁺(C₀) = C⁺',5:'S(R): ensured, если R ⊆ C⁻; not_ensured, если R ⊈ C⁺; undetermined в остальных случаях',6:'B = V ∖ C⁺; B ∩ C₀ = ∅; B ∩ R ≠ ∅; ∀ e ∈ E⁺: oₑ ∈ B ⇒ Iₑ ∩ B ≠ ∅',7:'L = ∑ₑ∈E |Iₑ|; T = O(|V| + |E| + L)',8:'Dᵣ₊₁ = Dᵣ ∪ {oₑ : e ∈ F, Iₑ ∩ Dᵣ ≠ ∅}; D = ⋃ᵣ≥₀ Dᵣ'}
xs_path=Path(os.environ.get('MML2OMML_XSL',r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL'))
transform=etree.XSLT(etree.parse(str(xs_path)))

doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
for side in ['top_margin','bottom_margin','left_margin','right_margin']:setattr(sec,side,Cm(2.54))
sec.header_distance=Cm(1.25);sec.footer_distance=Cm(1.25)
styles=doc.styles
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
 s=styles[name];s.font.name='Times New Roman';s.font.size=Pt(11);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.line_spacing=1;s.paragraph_format.space_after=Pt(0)
 for rf in s.element.findall('.//'+qn('w:rFonts')):
  for attr in list(rf.attrib):
   if 'theme' in attr.lower():del rf.attrib[attr]
 for tag in ['w:pBdr','w:spacing','w:kern']:
  if tag=='w:spacing':
   # Remove character tracking, retaining paragraph spacing.
   nodes=s.element.findall('./'+qn('w:rPr')+'/'+qn(tag))
  else:nodes=s.element.findall('.//'+qn(tag))
  for z in nodes:z.getparent().remove(z)
 for z in s.element.findall('.//'+qn('w:szCs')):z.set(qn('w:val'),'22')
styles['Normal'].paragraph_format.first_line_indent=Cm(.75)
styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
styles['Normal'].paragraph_format.widow_control=True
for name in ['Heading 1','Heading 2']:
 s=styles[name];s.font.bold=True;s.paragraph_format.space_before=Pt(8);s.paragraph_format.space_after=Pt(4);s.paragraph_format.first_line_indent=Cm(0);s.paragraph_format.keep_with_next=True;s.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
styles['Title'].font.bold=True;styles['Title'].paragraph_format.space_before=Pt(3);styles['Title'].paragraph_format.space_after=Pt(5)
styles['Title'].paragraph_format.first_line_indent=Cm(0);styles['Title'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
styles['Title'].paragraph_format.keep_with_next=True
styles['Caption'].font.italic=False
for prop in ['autoHyphenation']:
 z=OxmlElement('w:'+prop);z.set(qn('w:val'),'true');doc.settings.element.append(z)
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');styles['Normal'].element.get_or_add_rPr().append(lang)
mathPr=OxmlElement('m:mathPr');mf=OxmlElement('m:mathFont');mf.set(qn('m:val'),'Cambria Math');mathPr.append(mf);doc.settings.element.append(mathPr)

def math_node(content,display=False):
 m=etree.fromstring(('<math xmlns="http://www.w3.org/1998/Math/MathML">'+content+'</math>').encode());node=transform(m).getroot()
 for mr in node.findall('.//'+qn('m:r')):
  pr=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'22');pr.append(sz);mr.insert(1 if len(mr) and mr[0].tag==qn('m:rPr') else 0,pr)
 if display:
  for nary in node.findall('.//'+qn('m:naryPr')):
   loc=nary.find(qn('m:limLoc'))
   if loc is None:loc=OxmlElement('m:limLoc');nary.append(loc)
   loc.set(qn('m:val'),'undOvr')
 return node
def para(text='',align=None,bold=None,italic=None,indent=None,style=None,math_inline=False):
 p=doc.add_paragraph(style=style)
 if align is not None:p.alignment=align
 if indent is not None:p.paragraph_format.first_line_indent=Cm(indent)
 parts=re.split(r'(cl_E\(C₀\)|C₀|C⁻|C\+|E⁻|E⁺|Iₑ|oₑ|D₀|Jᴰ|Kᴰ|(?<!\w)[kEtvVRBDFℓ](?!\w))',text) if math_inline else [text]
 for part in parts:
  if part=='cl_E(C₀)':p._p.append(math_node(row(cl(E))));continue
  rrun=p.add_run(part)
  if bold is not None:rrun.bold=bold
  if italic is not None:rrun.italic=italic
  if math_inline and re.fullmatch(r'C₀|C⁻|C\+|E⁻|E⁺|Iₑ|oₑ|D₀|Jᴰ|Kᴰ|[kEtvVRBDFℓ]',part):rrun.italic=True
 if 'формулой (6)' in text:p.paragraph_format.keep_with_next=True
 if text.startswith('Никитин Пётр'):p.paragraph_format.keep_with_next=True
 return p
def caption(ru,en):
 p=para(ru,align=WD_ALIGN_PARAGRAPH.CENTER,indent=0);p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(5)
 q=para(en,align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,indent=0);q.paragraph_format.keep_with_next=True;q.paragraph_format.space_after=Pt(3)
def grid(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 pr=t._tbl.tblPr;borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  v=OxmlElement('w:'+edge);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'000000');borders.append(v)
 pr.append(borders)
 margin=OxmlElement('w:tblCellMar')
 for key,val in [('top',60),('bottom',60),('left',85),('right',85)]:
  v=OxmlElement('w:'+key);v.set(qn('w:w'),str(val));v.set(qn('w:type'),'dxa');margin.append(v)
 pr.append(margin)
 for i,(cell,text) in enumerate(zip(t.rows[0].cells,headers)):
  cell.text=text;cell.width=Cm(widths[i])
  for r in cell.paragraphs[0].runs:r.bold=True
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 for rowvals in rows:
  cs=t.add_row().cells
  for i,(c,txt) in enumerate(zip(cs,rowvals)):c.text=str(txt);c.width=Cm(widths[i])
 for rr in t.rows:
  no=OxmlElement('w:cantSplit');rr._tr.get_or_add_trPr().append(no)
  for i,c in enumerate(rr.cells):
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for sh in list(c._tc.get_or_add_tcPr().findall(qn('w:shd'))):sh.getparent().remove(sh)
   for p in c.paragraphs:
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(0)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs:r.font.size=Pt(11)
 return t
table_records={}
def add_table(kind):
 if kind=='classification':
  names={'clean':'Полная цепочка','time':'Предпосылка опаздывает','level':'Недостаточный уровень','diagnostic':'Отсутствует контроль','version':'Версия не согласована','and':'Нет одной предпосылки','cycle':'Цикл без внешней опоры','unknown':'Контроль не размечен','alternative':'Есть допустимая альтернатива'}
  st={'ensured':'Обеспечено','not_ensured':'Не обеспечено','undetermined':'Недостаточно сведений'}
  headers=['Тип примера','M','O','A','Полная модель'];rows=[[names[x['scenario']],x['matrix'],x['ordinary'],x['no_time_version'],st[x['truth']]] for x in R['classifier_summary']]
  caption('Таблица 1. Положительные ответы упрощений из 40 вариантов каждого типа','Table 1. Positive decisions of simplified models out of 40 variants per case type')
  grid(headers,rows,[6.2,1,1,1,6.72])
 elif kind=='timings':
  headers=['Вершины','Режим','|D|','Полный, мс','Обновление, мс'];rows=[]
  for n in [1000,5000,10000,20000]:
   for scope in ['local','global']:
    a=[x for x in R['timings'] if x['vertices']==n and x['change_scope']==scope]
    rows.append([n,'Локальный' if scope=='local' else 'Глобальный',int(statistics.median(x['affected_vertices'] for x in a)),f"{statistics.median(x['full_ms'] for x in a):.3f}".replace('.',','),f"{statistics.median(x['local_ms'] for x in a):.3f}".replace('.',',')])
  caption('Таблица 2. Медианное время пересчёта при готовом индексе','Table 2. Median recomputation times with a pre-existing index')
  grid(headers,rows,[2.3,3.4,2.1,3.2,4.92])
 elif kind=='courses':
  headers=['Дисциплина','Семестры','Коды из плана','№'];rows=[]
  for x in json.loads((B/'curriculum_fragment.json').read_text(encoding='utf8')):
   codes=x['competency_codes']
   if x['row']==54:codes='MF-2, MF-4; ML-2–ML-8'
   rows.append([x['name'],', '.join(map(str,x['semesters'])),codes,x['row']])
  caption('Таблица 3. Документальные сведения листа «Дисциплины» учебного плана','Table 3. Documentary facts from the curriculum worksheet “Дисциплины”')
  grid(headers,rows,[6.3,2.6,5.5,1.52])
 table_records[kind]={'headers':headers,'rows':rows}

def equation(n):
 p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
 p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(7.5),WD_TAB_ALIGNMENT.CENTER)
 p.paragraph_format.tab_stops.add_tab_stop(Cm(15.92),WD_TAB_ALIGNMENT.RIGHT)
 p.add_run('\t');node=math_node(eqs[n],display=True)
 p._p.append(node);p.add_run('\t('+str(n)+')')

source=(B/'manuscript.md').read_text(encoding='utf8')
for line in source.split('\n\n'):
 line=line.strip()
 if not line:continue
 if line.startswith('@@EQ:'):equation(int(re.search(r'\d+',line).group()));continue
 if line.startswith('@@TABLE:'):add_table(line.split(':')[1].split('@')[0]);continue
 if line=='@@FIG:pipeline@@':
  p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  p.add_run().add_picture(str(B/'figure_1_pipeline.png'),width=Cm(15.5))
  caption('Рис. 1. Проверка обеспеченности и уточнение кандидатных связей','Fig. 1. Provision verification and refinement of candidate annotations')
  # Following body may start on the next page; do not chain figure captions to it.
  doc.paragraphs[-1].paragraph_format.keep_with_next=False
  continue
 if line=='@@REFERENCES@@':
  for i,(text,doi) in enumerate(refs,1):
   p=para(str(i)+'. '+text+(' DOI: '+doi+'.' if doi else ''),indent=0)
   p.paragraph_format.left_indent=Cm(.6);p.paragraph_format.first_line_indent=Cm(-.6)
  continue
 if line.startswith('# '):para(line[2:],style='Title');continue
 if line.startswith('## '):
  text=line[3:]
  if text.startswith('A Model'):para(text,style='Title')
  else:para(text,style='Heading 1')
  continue
 if line in ['П. В. Никитин','Petr V. Nikitin']:para(line,align=WD_ALIGN_PARAGRAPH.CENTER,bold=True,indent=0);continue
 if line.startswith('Финансовый университет') or line.startswith('Financial University'):
  para(line,align=WD_ALIGN_PARAGRAPH.CENTER,italic=True,indent=0);continue
 if line.startswith('УДК'):para(line,align=WD_ALIGN_PARAGRAPH.LEFT,indent=0);continue
 para(line,math_inline=not (line.startswith('Petr V.') or line.startswith('Никитин Пётр')))

# A small conventional page number, with no decorative header.
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');p._p.append(f)
doc.core_properties.title='Модель и алгоритм проверки структурной согласованности междисциплинарной образовательной программы'
doc.core_properties.author='Пётр Владимирович Никитин'
doc.core_properties.subject='Структурная проверка программы по специальности 1.2.2'
path=OUT/NAME;doc.save(path)
alltext='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for rr in t.rows for c in rr.cells)+'\n'+'\n'.join(eq_plain.values())
metrics={'docx':str(path),'characters_including_tables_and_equations':len(alltext),'figure_equivalent_characters':round(6700*(155*53.9)/(190*270)),'english_abstract_words':len(source.split('Abstract. ')[1].split('\n\nKeywords')[0].split()),'references':len(refs),'native_equations':len(eqs),'tables':len(doc.tables),'em_dashes':alltext.count('—'),'processor':M.get('processor'),'timing_runs':len(R['timings'])*7}
metrics['estimated_total_with_figure']=metrics['characters_including_tables_and_equations']+metrics['figure_equivalent_characters']
(B/'document_metrics.json').write_text(json.dumps(metrics,ensure_ascii=False,indent=2),encoding='utf8')
(B/'article_tables.json').write_text(json.dumps(table_records,ensure_ascii=False,indent=2),encoding='utf8')
(B/'references.json').write_text(json.dumps([{'number':i,'text':t,'doi':d} for i,(t,d) in enumerate(refs,1)],ensure_ascii=False,indent=2),encoding='utf8')
readable=source
for n,plain in eq_plain.items():readable=readable.replace(f'@@EQ:{n}@@',plain+f'  ({n})')
for kind,tabledata in table_records.items():
 head=tabledata['headers'];body=tabledata['rows'];md='| '+' | '.join(head)+' |\n| '+' | '.join('---' for _ in head)+' |\n'+'\n'.join('| '+' | '.join(map(str,rr))+' |' for rr in body)
 readable=readable.replace('@@TABLE:'+kind+'@@',md)
readable=readable.replace('@@FIG:pipeline@@','![Проверка обеспеченности](figure_1_pipeline.png)')
readable=readable.replace('@@REFERENCES@@','\n\n'.join(f'{i}. {t}'+(' DOI: '+d if d else '') for i,(t,d) in enumerate(refs,1)))
(B/'article_readable.md').write_text(readable,encoding='utf8')
print(json.dumps(metrics,ensure_ascii=False,indent=2))
